# -*- coding: utf-8 -*-
"""
Created on Wed Mar 11 11:33:34 2026 alpha 9 project @author: Maj0131
"""
import os
os.environ['KMP_DUPLICATE_LIB_OK'] = 'True'

import csv
import pandas as pd
import numpy as np
import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, ttk, messagebox, font
import sys
import torch
import torch.nn as nn
import threading
import matplotlib
matplotlib.use("TkAgg") 
import matplotlib.pyplot as plt
import pickle 
from itertools import product
from sklearn.metrics import mean_squared_error, r2_score, confusion_matrix, ConfusionMatrixDisplay
from sklearn.svm import SVC, SVR  # SVM support
from sklearn.preprocessing import StandardScaler, MinMaxScaler
from sklearn.linear_model import LinearRegression, LogisticRegression
from sklearn.tree import DecisionTreeClassifier, DecisionTreeRegressor
from sklearn.ensemble import (
    RandomForestClassifier, RandomForestRegressor, 
    HistGradientBoostingClassifier, HistGradientBoostingRegressor
    )
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageTk, ImageDraw
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

# --- Model & Hyperparameter Dictionary ---
MODEL_LIBRARY = {
    "ANN (MLP)": {
       "Hidden Layers": {"type": "int", "default": "3"},
       "Neurons": {"type": "int", "default": "64"},
       "Dist. Mode": {"type": "list", "options": ["Constant", "Pyramid", "Custom List"]},
       "Batch Size": {"type": "int", "default": "32"},
       "Epochs": {"type": "int", "default": "50"},        
       "Activation": {"type": "list", "options": ["ReLU", "Sigmoid", "Tanh", "LeakyReLU"]},
       "Optimizer": {"type": "list", "options": ["Adam", "SGD", "RMSprop"]},
       "Learning Rate": {"type": "float", "default": "0.001"}
       },
    "RBF Network": {
        "Num Centers": {"type": "int", "default": "20"},
        "Sigma": {"type": "float", "default": "1.0"},
        "Batch Size": {"type": "int", "default": "32"},
        "Epochs": {"type": "int", "default": "100"},
        "Learning Rate": {"type": "float", "default": "0.01"}
     },
    "CNN": {
        "Filters": {"type": "int", "default": "32"},
        "Kernel Size": {"type": "int", "default": "3"},
        "Pool Size": {"type": "int", "default": "2"},
        "Batch Size": {"type": "int", "default": "32"},
        "Epochs": {"type": "int", "default": "50"},
        "Learning Rate": {"type": "float", "default": "0.001"}
    },
    "LSTM": {
         "Hidden Dim": {"type": "int", "default": "128"},
         "Num Layers": {"type": "int", "default": "2"},
         "Dropout": {"type": "float", "default": "0.2"},
         "Batch Size": {"type": "int", "default": "32"},
         "Epochs": {"type": "int", "default": "50"},
         "Learning Rate": {"type": "float", "default": "0.001"}
     },
    "SVM": {
          "Kernel": {"type": "list", "options": ["rbf", "linear", "poly", "sigmoid"]},
          "C (Penalty)": {"type": "float", "default": "1.0"},
          "Gamma": {"type": "list", "options": ["scale", "auto"]}
     },
    "LR (Linear/Logistic)": {
        "Penalty": {"type": "list", "options": ["l2", "l1", "none"]},
        "Max Iter": {"type": "int", "default": "1000"},
        "Solver": {"type": "list", "options": ["lbfgs", "liblinear", "saga"]}
    },
    "Decision Tree": {
        "Max Depth": {"type": "int", "default": "10"},
        "Min Samples Split": {"type": "int", "default": "2"}
    },
    "Random Forest": {
        "n_Estimators": {"type": "int", "default": "100"},
        "Max Depth": {"type": "int", "default": "10"}
    },
    "Gradient Boosting": {
        "Learning Rate": {"type": "float", "default": "0.1"},
        "Max Iter": {"type": "int", "default": "100"},
        "Max Depth": {"type": "int", "default": "5"}
    },
       
}
# ---------- FEATURE EXTRACTION ----------
def sliding_window_features(df, inputs, target, win_size, overlap, feats):
    return extract_features(df, inputs, target, win_size, overlap, feats, mode="sliding")

def fourier_window_features(df, inputs, target, win_size, overlap, feats):
    return extract_features(df, inputs, target, win_size, overlap, feats, mode="fourier")

def frequency_domain_features(df, inputs, target, win_size, overlap, feats):
    return extract_features(df, inputs, target, win_size, overlap, feats, mode="frequency")

def normalize_numeric_series(series):
    if series is None:
        return series
    s = series.copy()
    # First try standard pandas conversion
    try:
        return pd.to_numeric(s, errors="raise")
    except Exception:
        pass

    def clean_value(value):
        if value is None:
            return value
        text = str(value).strip()
        if text == "":
            return value
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        return text

    cleaned = s.map(clean_value)
    converted = pd.to_numeric(cleaned, errors="coerce")
    if converted.notna().sum() >= max(1, int(len(converted) * 0.8)):
        return converted
    return s

def extract_features(df, inputs, target, win_size, overlap, feats, mode="sliding"):
    step = max(1, int(win_size * (1 - overlap / 100)))
    X, y = [], []
    feat_funcs = {
        "Mean": np.mean,
        "RMS": lambda x: np.sqrt(np.mean(x**2)),
        "Std Dev": np.std,
        "Min": np.min,
        "Max": np.max,
        "Peak-to-Peak": lambda x: np.ptp(x)
    }
    selected = {k: v for k, v in feat_funcs.items() if feats.get(k, False)}
    for start in range(0, len(df) - win_size + 1, step):
        w = df.iloc[start:start + win_size]
        feat_vec = []

        for col in inputs:
            data = normalize_numeric_series(w[col]).astype(float).to_numpy()
            if mode == "sliding":
                values = data
            else:
                transformed = np.fft.rfft(data - np.mean(data))
                values = np.abs(transformed)
            if mode == "frequency":
                magnitude = values
                if np.sum(magnitude) > 0:
                    freqs = np.arange(len(magnitude))
                    centroid = np.sum(freqs * magnitude) / np.sum(magnitude)
                    bandwidth = np.sqrt(np.sum(((freqs - centroid) ** 2) * magnitude) / np.sum(magnitude))
                else:
                    centroid = 0.0
                    bandwidth = 0.0
                feature_values = {
                    "Mean": np.mean(magnitude),
                    "RMS": np.sqrt(np.mean(magnitude ** 2)),
                    "Std Dev": np.std(magnitude),
                    "Min": np.min(magnitude),
                    "Max": np.max(magnitude),
                    "Peak-to-Peak": np.ptp(magnitude),
                    "Centroid": centroid,
                    "Bandwidth": bandwidth
                }
                if feats.get("Mean", False): feat_vec.append(feature_values["Mean"])
                if feats.get("RMS", False): feat_vec.append(feature_values["RMS"])
                if feats.get("Std Dev", False): feat_vec.append(feature_values["Std Dev"])
                if feats.get("Min", False): feat_vec.append(feature_values["Min"])
                if feats.get("Max", False): feat_vec.append(feature_values["Max"])
                if feats.get("Peak-to-Peak", False): feat_vec.append(feature_values["Peak-to-Peak"])
                if feats.get("Centroid", False): feat_vec.append(feature_values["Centroid"])
                if feats.get("Bandwidth", False): feat_vec.append(feature_values["Bandwidth"])
            else:
                for f in selected.values():
                    feat_vec.append(f(values))

        X.append(feat_vec)
        y.append(w[target].iloc[-1])
    return np.array(X), np.array(y)
# ---------- TRAIN / VAL / TEST SPLIT ----------
def split_data(X, y, train_pct):
    n = len(X)
    idx = np.random.permutation(n)
    t = int(n * train_pct / 100)
    v = int((n - t) / 2)
    return (
        X[idx[:t]], y[idx[:t]],
        X[idx[t:t+v]], y[idx[t:t+v]],
        X[idx[t+v:]], y[idx[t+v:]]
    )
# ---------- SVM TRAINING (NEW) ----------
def train_svm(Xtr, ytr, params, target_type, progress_callback=None):
    kernel = params.get("Kernel", "rbf")
    c_val = float(params.get("C (Penalty)", 1.0))
    gamma = params.get("Gamma", "scale") 
    if target_type == "Regression":
        model = SVR(kernel=kernel, C=c_val, gamma=gamma)
    else:
        model = SVC(kernel=kernel, C=c_val, gamma=gamma, probability=True)
    print(f">> Training SVM ({target_type})...")
    if progress_callback:
        progress_callback(0, 1, "Training SVM...")
    model.fit(Xtr, ytr)
    if progress_callback:
        progress_callback(1, 1, "SVM training complete")
    return model
def train_cnn(Xtr, ytr, Xv, yv, params, num_outputs, target_type, stop_event=None, history=None, Xt=None, yt=None, progress_callback=None):
    device = "cuda" if torch.cuda.is_available() else "cpu"
    num_out = num_outputs   
    # Initialization
    model = CNN1D(
        in_channels=Xtr.shape[2], 
        filters=int(params["Filters"]),
        kernel_size=int(params["Kernel Size"]),
        pool_size=int(params["Pool Size"]),
        num_outputs=num_out
    ).to(device)
    loss_fn = nn.MSELoss() if target_type == "Regression" else nn.CrossEntropyLoss()
    y_dtype = torch.float32 if target_type == "Regression" else torch.long    
    optim = torch.optim.Adam(model.parameters(), lr=float(params.get("Learning Rate", 0.001)))
    
    Xtr_tensor = torch.tensor(Xtr, dtype=torch.float32)
    Xv_tensor = torch.tensor(Xv, dtype=torch.float32).to(device) if Xv is not None else None
    Xt_tensor = torch.tensor(Xt, dtype=torch.float32).to(device) if Xt is not None else None
    yv_tensor = torch.tensor(yv, dtype=torch.float32 if target_type == "Regression" else torch.long).to(device) if yv is not None else None
    yt_tensor = torch.tensor(yt, dtype=torch.float32 if target_type == "Regression" else torch.long).to(device) if yt is not None else None
    loader = torch.utils.data.DataLoader(
        torch.utils.data.TensorDataset(torch.tensor(Xtr, dtype=torch.float32), 
                                       torch.tensor(ytr, dtype=y_dtype)),
        batch_size=int(params.get("Batch Size", 32)), shuffle=True
    )
    epochs = int(params.get("Epochs", 50))
    print(f">> Starting CNN Training ({target_type})...")
    if progress_callback:
        progress_callback(0, epochs, f"CNN Epoch 0/{epochs}")

    def get_nn_preds(data):
        if data is None or len(data) == 0:
            return np.array([])
        model.eval()
        with torch.no_grad():
            t_data = torch.tensor(data, dtype=torch.float32, device=device)
            raw = model(t_data)
            if target_type == "Regression":
                preds = raw.cpu().numpy().reshape(-1)
            else:
                preds = torch.argmax(raw, dim=1).cpu().numpy()
        model.train()
        return preds

    for epoch in range(1, epochs + 1):
        if stop_event is not None and stop_event.is_set():
            print(">> CNN training canceled.")
            break
        model.train()
        total_loss = 0
        for xb, yb in loader:
            xb, yb = xb.to(device), yb.to(device)
            pred = model(xb)
            if target_type == "Regression":
                pred = pred.reshape(-1)
                yb = yb.reshape(-1)
            loss = loss_fn(pred, yb)
            optim.zero_grad()
            loss.backward()
            optim.step()
            total_loss += loss.item()
        if history is not None:
            if target_type == "Regression":
                train_pred = get_nn_preds(Xtr)
                train_loss = float(
                    loss_fn(
                        torch.tensor(train_pred, dtype=torch.float32, device=device).reshape(-1),
                        torch.tensor(ytr, dtype=torch.float32, device=device).reshape(-1)
                    )
                ) if len(train_pred) > 0 else None
                val_loss = None
                test_loss = None
                if Xv is not None and yv is not None:
                    val_pred = get_nn_preds(Xv)
                    val_loss = float(
                        loss_fn(
                            torch.tensor(val_pred, dtype=torch.float32, device=device).reshape(-1),
                            torch.tensor(yv, dtype=torch.float32, device=device).reshape(-1)
                        )
                    ) if len(val_pred) > 0 else None
                if Xt is not None and yt is not None:
                    test_pred = get_nn_preds(Xt)
                    test_loss = float(
                        loss_fn(
                            torch.tensor(test_pred, dtype=torch.float32, device=device).reshape(-1),
                            torch.tensor(yt, dtype=torch.float32, device=device).reshape(-1)
                        )
                    ) if len(test_pred) > 0 else None
                history["train"].append(train_loss)
                history["val"].append(val_loss)
                history["test"].append(test_loss)
            else:
                train_acc = float(np.mean(get_nn_preds(Xtr) == ytr)) if len(Xtr) > 0 else None
                val_acc = float(np.mean(get_nn_preds(Xv) == yv)) if Xv is not None and len(Xv) > 0 else None
                test_acc = float(np.mean(get_nn_preds(Xt) == yt)) if Xt is not None and len(Xt) > 0 else None
                history["train"].append(train_acc)
                history["val"].append(val_acc)
                history["test"].append(test_acc)
        if progress_callback:
            progress_callback(epoch, epochs, f"CNN Epoch {epoch}/{epochs}")
        if epoch % 5 == 0 or epoch == 1:
            if Xv is not None:
                print(f"CNN Epoch {epoch:03d}/{epochs} | Loss: {total_loss/len(loader):.4f}")
            else:
                print(f"CNN Epoch {epoch:03d}/{epochs} | Loss: {total_loss/len(loader):.4f}")
    if progress_callback and not (stop_event is not None and stop_event.is_set()):
        progress_callback(epochs, epochs, "CNN training complete")
    print(">> CNN Training Complete!")
    return model
#------------trees 0.6.9.1-------------
def train_tree_methods(Xtr, ytr, params, target_type, model_type, progress_callback=None):
    depth = int(params.get("Max Depth", 10))
    if depth == 0: depth = None 
    
    if model_type == "Decision Tree":
        min_split = int(params.get("Min Samples Split", 2))
        if target_type == "Regression":
            model = DecisionTreeRegressor(max_depth=depth, min_samples_split=min_split)
        else:
            model = DecisionTreeClassifier(max_depth=depth, min_samples_split=min_split)            
    elif model_type == "Random Forest":
        n_est = int(params.get("n_Estimators", 100))
        if target_type == "Regression":
            model = RandomForestRegressor(n_estimators=n_est, max_depth=depth, n_jobs=-1)
        else:
            model = RandomForestClassifier(n_estimators=n_est, max_depth=depth, n_jobs=-1)
    elif model_type == "Gradient Boosting":
        lr = float(params.get("Learning Rate", 0.1))
        iters = int(params.get("Max Iter", 100))
        if target_type == "Regression":
            model = HistGradientBoostingRegressor(learning_rate=lr, max_iter=iters, max_depth=depth)
        else:
            model = HistGradientBoostingClassifier(learning_rate=lr, max_iter=iters, max_depth=depth)            
    print(f">> Training {model_type}...")
    if progress_callback:
        progress_callback(0, 1, f"Training {model_type}...")
    model.fit(Xtr, ytr)
    if progress_callback:
        progress_callback(1, 1, f"{model_type} training complete")
    print(f">> {model_type} Training Complete!")
    return model
#------------Regression 0.6.5.1-------------
def train_lr(Xtr, ytr, params, target_type, progress_callback=None):
    # Mapping GUI strings to sklearn-friendly values
    pen = params.get("Penalty", "l2").lower()
    if pen == "none":
        pen = None       
    max_i = int(params.get("Max Iter", 1000))
    solv = params.get("Solver", "lbfgs").lower()
    if target_type == "Regression":
        print(">> Training Linear Regression (OLS)...")
        model = LinearRegression()
    else:
        print(f">> Training Logistic Regression | Solver: {solv} | Penalty: {pen}")
        if solv == "liblinear" and pen is None:
            print("   ! liblinear doesn't support None penalty. Falling back to l2.")
            pen = "l2"            
        model = LogisticRegression(
            penalty=pen, 
            max_iter=max_i, 
            solver=solv, 
            multi_class='auto'
        )    
    if progress_callback:
        progress_callback(0, 1, "Training LR model...")
    model.fit(Xtr, ytr)
    if progress_callback:
        progress_callback(1, 1, "LR training complete")
    print(">> LR Training Complete!")
    return model
# ---------- CNN ----------
class CNN1D(nn.Module):
    def __init__(self, in_channels, filters, kernel_size, pool_size, num_outputs):
        super().__init__()  
        self.conv_block = nn.Sequential(
            nn.Conv1d(in_channels, filters, kernel_size=kernel_size, padding=1),
            nn.ReLU(),
            nn.MaxPool1d(pool_size),
            nn.Conv1d(filters, filters*2, kernel_size=kernel_size, padding=1),
            nn.ReLU(),
            nn.AdaptiveAvgPool1d(1) 
        )
        self.fc = nn.Linear(filters*2, num_outputs)
    def forward(self, x):    
        x = x.transpose(1, 2)
        x = self.conv_block(x)
        x = x.view(x.size(0), -1) # Flatten
        return self.fc(x)
# ---------- RBF NETWORK ----------
class RBFNet(nn.Module):
    def __init__(self, in_dim, num_centers, sigma, out_dim):
        super(RBFNet, self).__init__()
        self.num_centers = num_centers
        self.sigma = sigma        
        # Centers are parameters that the network will "learn" or position
        self.centers = nn.Parameter(torch.randn(num_centers, in_dim))
        self.linear = nn.Linear(num_centers, out_dim)
    def kernel_fun(self, batches):
        n_input = batches.size(0)
        dist = (batches.unsqueeze(1) - self.centers.unsqueeze(0)).pow(2).sum(-1)
        # Gaussian: exp(-dist / (2 * sigma^2))
        return torch.exp(-dist / (2 * self.sigma**2))
    def forward(self, x):
        radial_val = self.kernel_fun(x)
        return self.linear(radial_val)
def train_rbf(Xtr, ytr, Xv, yv, params, num_outputs, target_type, stop_event=None, history=None, Xt=None, yt=None, progress_callback=None):
    device = "cuda" if torch.cuda.is_available() else "cpu"
    num_out = num_outputs    
    model = RBFNet(Xtr.shape[1], int(params["Num Centers"]), 
                   float(params["Sigma"]), num_out).to(device)    
    loss_fn = nn.MSELoss() if target_type == "Regression" else nn.CrossEntropyLoss()
    y_t = torch.tensor(ytr, dtype=torch.float32 if target_type == "Regression" else torch.long).to(device)
    x_t = torch.tensor(Xtr, dtype=torch.float32).to(device)
    Xt_tensor = torch.tensor(Xt, dtype=torch.float32).to(device) if Xt is not None else None
    yt_tensor = torch.tensor(yt, dtype=torch.float32 if target_type == "Regression" else torch.long).to(device) if yt is not None else None
    optim = torch.optim.Adam(model.parameters(), lr=float(params.get("Learning Rate", 0.01))) 
    epochs = int(params.get("Epochs", 50))
    print(f">> Starting RBF Training for {epochs} epochs...") # Checkpoint 1
    if progress_callback:
        progress_callback(0, epochs, f"RBF Epoch 0/{epochs}")

    def get_nn_preds(data):
        was_training = model.training
        model.eval()
        with torch.no_grad():
            t_data = torch.tensor(data, dtype=torch.float32).to(device)
            raw = model(t_data)
        if was_training:
            model.train()
        if target_type == "Regression":
            return raw.cpu().numpy().flatten()
        return torch.argmax(raw, dim=1).cpu().numpy()

    for epoch in range(1, epochs + 1):
        if stop_event is not None and stop_event.is_set():
            print(">> RBF training canceled.")
            break
        model.train()
        optim.zero_grad()
        pred = model(x_t)
        
        loss = loss_fn(pred.squeeze() if target_type == "Regression" else pred, y_t)
        loss.backward()
        optim.step()
        if history is not None:
            if target_type == "Regression":
                train_loss = loss.item()
                val_loss = None
                test_loss = None
                if Xv is not None and yv is not None:
                    val_pred = get_nn_preds(Xv)
                    val_loss = float(loss_fn(torch.tensor(val_pred, dtype=torch.float32), torch.tensor(yv, dtype=torch.float32))) if isinstance(loss_fn, nn.MSELoss) else None
                if Xt is not None and yt is not None:
                    test_pred = get_nn_preds(Xt)
                    test_loss = float(loss_fn(torch.tensor(test_pred, dtype=torch.float32), torch.tensor(yt, dtype=torch.float32))) if isinstance(loss_fn, nn.MSELoss) else None
                history["train"].append(train_loss)
                history["val"].append(val_loss)
                history["test"].append(test_loss)
            else:
                train_acc = np.mean(get_nn_preds(Xtr) == ytr)
                val_acc = np.mean(get_nn_preds(Xv) == yv) if Xv is not None else None
                test_acc = np.mean(get_nn_preds(Xt) == yt) if Xt is not None else None
                history["train"].append(train_acc)
                history["val"].append(val_acc)
                history["test"].append(test_acc)
        if progress_callback:
            progress_callback(epoch, epochs, f"RBF Epoch {epoch}/{epochs}")
        if epoch % 10 == 0 or epoch == 1:
            print(f"RBF Epoch {epoch:03d}/{epochs} | Loss: {loss.item():.4f}")
    if progress_callback and not (stop_event is not None and stop_event.is_set()):
        progress_callback(epochs, epochs, "RBF training complete")
    print(">> RBF Training Complete!") # Checkpoint 2
    return model
# ---------- LSTM-------------------------
class LSTMModel(nn.Module):
    def __init__(self, in_dim, hidden_dim, num_layers, dropout, num_outputs=1):
        super().__init__()
        self.hidden_dim = hidden_dim
        self.num_layers = num_layers        
        # LSTM Layer
        self.lstm = nn.LSTM(in_dim, hidden_dim, num_layers, 
                            batch_first=True, dropout=dropout if num_layers > 1 else 0)        
        # Fully connected output layer
        self.fc = nn.Linear(hidden_dim, num_outputs)
    def forward(self, x):
        # x shape: (Batch, Seq_Len, Features)
        out, _ = self.lstm(x)
        # We only care about the output of the last time step
        out = self.fc(out[:, -1, :])
        return out
def train_lstm(Xtr, ytr, Xv, yv, params, num_outputs, target_type, stop_event=None, history=None, Xt=None, yt=None, progress_callback=None):
    device = "cuda" if torch.cuda.is_available() else "cpu"   
    num_out = num_outputs
    # Task Logic
    if target_type == "Regression":
        loss_fn = nn.MSELoss()
        ytr_tensor = torch.tensor(ytr, dtype=torch.float32)
        yv_tensor  = torch.tensor(yv,  dtype=torch.float32).to(device) if yv is not None else None
    else:
        loss_fn = nn.CrossEntropyLoss()
        ytr_tensor = torch.tensor(ytr, dtype=torch.long)
        yv_tensor  = torch.tensor(yv, dtype=torch.long).to(device) if yv is not None else None
    Xtr_tensor = torch.tensor(Xtr, dtype=torch.float32)
    Xv_tensor  = torch.tensor(Xv,  dtype=torch.float32).to(device) if Xv is not None else None
    Xt_tensor  = torch.tensor(Xt, dtype=torch.float32).to(device) if Xt is not None else None
    yt_tensor  = torch.tensor(yt, dtype=torch.float32 if target_type == "Regression" else torch.long).to(device) if yt is not None else None
    model = LSTMModel(
        in_dim=Xtr.shape[2], # Number of features
        hidden_dim=int(params["Hidden Dim"]),
        num_layers=int(params["Num Layers"]),
        dropout=float(params["Dropout"]),
        num_outputs=num_out
    ).to(device)
    lr = float(params.get("Learning Rate", 0.001))
    optim = torch.optim.Adam(model.parameters(), lr=lr)
    loader = torch.utils.data.DataLoader(
        torch.utils.data.TensorDataset(Xtr_tensor, ytr_tensor), 
        batch_size=int(params.get("Batch Size", 32)), shuffle=True
    )
    epochs = int(params.get("Epochs", 50))
    if progress_callback:
        progress_callback(0, epochs, f"LSTM Epoch 0/{epochs}")

    def get_nn_preds(data):
        was_training = model.training
        model.eval()
        with torch.no_grad():
            t_data = torch.tensor(data, dtype=torch.float32).to(device)
            raw = model(t_data)
        if was_training:
            model.train()
        if target_type == "Regression":
            return raw.cpu().numpy().flatten()
        return torch.argmax(raw, dim=1).cpu().numpy()

    for epoch in range(1, epochs + 1): 
        if stop_event is not None and stop_event.is_set():
            print(">> LSTM training canceled.")
            break
        model.train()
        total_loss = 0.0
        for xb, yb in loader:
            xb, yb = xb.to(device), yb.to(device)
            pred = model(xb)            
            # Regression vs Classification loss handling
            loss = loss_fn(pred.squeeze() if target_type == "Regression" else pred, yb)            
            optim.zero_grad()
            loss.backward()
            optim.step()
            total_loss += loss.item()
        if history is not None:
            if target_type == "Regression":
                train_loss = total_loss / len(loader)
                val_loss = None
                test_loss = None
                if Xv_tensor is not None and yv_tensor is not None:
                    with torch.no_grad():
                        v_pred = model(Xv_tensor)
                        vloss = loss_fn(v_pred.squeeze(), yv_tensor)
                        val_loss = vloss.item()
                if Xt_tensor is not None and yt_tensor is not None:
                    with torch.no_grad():
                        t_pred = model(Xt_tensor)
                        tloss = loss_fn(t_pred.squeeze(), yt_tensor)
                        test_loss = tloss.item()
                history["train"].append(train_loss)
                history["val"].append(val_loss)
                history["test"].append(test_loss)
            else:
                train_acc = np.mean(get_nn_preds(Xtr) == ytr)
                val_acc = np.mean(get_nn_preds(Xv) == yv) if Xv is not None else None
                test_acc = np.mean(get_nn_preds(Xt) == yt) if Xt is not None else None
                history["train"].append(train_acc)
                history["val"].append(val_acc)
                history["test"].append(test_acc)
        if progress_callback:
            progress_callback(epoch, epochs, f"LSTM Epoch {epoch}/{epochs}")
        # Optional: Add a print statement to see progress in your console
        if epoch % 5 == 0 or epoch == 1:
            print(f"LSTM Epoch {epoch:03d}/{epochs} | Loss: {total_loss/len(loader):.4f}")
    if progress_callback and not (stop_event is not None and stop_event.is_set()):
        progress_callback(epochs, epochs, "LSTM training complete")
    print(">> LTSM Training Complete!")       
    return model
# ---------- PYTORCH MLP (MODIFIED FOR CLASSIFICATION) ----------
class MLP(nn.Module):
    def __init__(self, in_dim, hidden_layers, neurons, activation, num_outputs=1):
        super().__init__()
        act_map = {
            "ReLU": nn.ReLU, "Sigmoid": nn.Sigmoid,
            "Tanh": nn.Tanh, "LeakyReLU": nn.LeakyReLU
        }
        layers = []
        prev = in_dim
        for _ in range(hidden_layers):
            layers.append(nn.Linear(prev, neurons))
            layers.append(act_map[activation]())
            prev = neurons
        
        layers.append(nn.Linear(prev, num_outputs)) # Dynamic output neurons
        self.net = nn.Sequential(*layers)

    def forward(self, x):
        return self.net(x)
# ---------- TRAINING LOOP (MODIFIED FOR DUAL MODE) ----------
def train_mlp(Xtr, ytr, Xv, yv, params, num_outputs, target_type, stop_event=None, history=None, Xt=None, yt=None, progress_callback=None):
    device = "cuda" if torch.cuda.is_available() else "cpu"
    num_out = num_outputs
    # --- Task Logic Selection ---
    if target_type == "Regression":
        loss_fn = nn.MSELoss()
        ytr_tensor = torch.tensor(ytr, dtype=torch.float32)
        yv_tensor  = torch.tensor(yv, dtype=torch.float32).to(device) if yv is not None else None
    else:
        # For Binary or Multi-Class
        loss_fn = nn.CrossEntropyLoss()
        ytr_tensor = torch.tensor(ytr, dtype=torch.long)
        yv_tensor  = torch.tensor(yv, dtype=torch.long).to(device) if yv is not None else None
    Xtr_tensor = torch.tensor(Xtr, dtype=torch.float32)
    Xv_tensor  = torch.tensor(Xv,  dtype=torch.float32).to(device) if Xv is not None else None
    Xt_tensor  = torch.tensor(Xt, dtype=torch.float32).to(device) if Xt is not None else None
    yt_tensor  = torch.tensor(yt, dtype=torch.float32 if target_type == "Regression" else torch.long).to(device) if yt is not None else None
    batch_sz = int(params.get("Batch Size", 32))
    epochs = int(params.get("Epochs", 50))
    lr = float(params.get("Learning Rate", 0.001))
    if progress_callback:
        progress_callback(0, epochs, f"MLP Epoch 0/{epochs}")
    dataset = torch.utils.data.TensorDataset(Xtr_tensor, ytr_tensor)
    loader = torch.utils.data.DataLoader(dataset, batch_size=batch_sz, shuffle=True)
    model = MLP(
        in_dim=Xtr.shape[1],
        hidden_layers=int(params["Hidden Layers"]),
        neurons=int(params["Neurons"]),
        activation=params["Activation"],
        num_outputs=num_out
    ).to(device)
    opt_choice = params.get("Optimizer", "Adam")
    if opt_choice == "Adam": optim = torch.optim.Adam(model.parameters(), lr=lr)
    elif opt_choice == "SGD": optim = torch.optim.SGD(model.parameters(), lr=lr)
    else: optim = torch.optim.RMSprop(model.parameters(), lr=lr)
    print(f">> Training Mode: {target_type} | Device: {device} | Outputs: {num_out}")

    def get_nn_preds(data):
        was_training = model.training
        model.eval()
        with torch.no_grad():
            t_data = torch.tensor(data, dtype=torch.float32).to(device)
            raw = model(t_data)
        if was_training:
            model.train()
        if target_type == "Regression":
            return raw.cpu().numpy().flatten()
        return torch.argmax(raw, dim=1).cpu().numpy()

    for epoch in range(1, epochs + 1):
        if stop_event is not None and stop_event.is_set():
            print(">> MLP training canceled.")
            break
        model.train()
        total_loss = 0.0
        for xb, yb in loader:
            xb, yb = xb.to(device), yb.to(device)
            pred = model(xb)
            
            # Regression expects (N), Classification expects (N, C)
            loss = loss_fn(pred.squeeze() if target_type == "Regression" else pred, yb)

            optim.zero_grad()
            loss.backward()
            optim.step()
            total_loss += loss.item()

        if history is not None:
            if target_type == "Regression":
                train_loss = total_loss / len(loader)
                val_loss = None
                test_loss = None
                if Xv_tensor is not None and yv_tensor is not None:
                    with torch.no_grad():
                        v_pred = model(Xv_tensor)
                        vloss = loss_fn(v_pred.squeeze(), yv_tensor)
                        val_loss = vloss.item()
                if Xt_tensor is not None and yt_tensor is not None:
                    with torch.no_grad():
                        t_pred = model(Xt_tensor)
                        tloss = loss_fn(t_pred.squeeze(), yt_tensor)
                        test_loss = tloss.item()
                history["train"].append(train_loss)
                history["val"].append(val_loss)
                history["test"].append(test_loss)
            else:
                train_acc = np.mean(get_nn_preds(Xtr) == ytr)
                val_acc = np.mean(get_nn_preds(Xv) == yv) if Xv is not None else None
                test_acc = np.mean(get_nn_preds(Xt) == yt) if Xt is not None else None
                history["train"].append(train_acc)
                history["val"].append(val_acc)
                history["test"].append(test_acc)

        if progress_callback:
            progress_callback(epoch, epochs, f"MLP Epoch {epoch}/{epochs}")

        model.eval()
        if Xv_tensor is not None and yv_tensor is not None:
            with torch.no_grad():
                v_pred = model(Xv_tensor)
                vloss = loss_fn(v_pred.squeeze() if target_type == "Regression" else v_pred, yv_tensor)
        else:
            vloss = None

        if epoch % 5 == 0 or epoch == 1:
            if vloss is not None:
                print(f"Epoch {epoch:03d}/{epochs} | Train Loss: {total_loss/len(loader):.4f} | Val Loss: {vloss.item():.4f}")
            else:
                print(f"Epoch {epoch:03d}/{epochs} | Train Loss: {total_loss/len(loader):.4f}")
    if progress_callback and not (stop_event is not None and stop_event.is_set()):
        progress_callback(epochs, epochs, "MLP training complete")
    print(">> MLP Training Complete!")
    return model
class ConsoleRedirect:
    def __init__(self, widget): self.widget = widget
    def write(self, msg): self.widget.after(0, self._write, msg)
    def _write(self, msg):
        self.widget.insert("end", msg)
        self.widget.see("end")
    def flush(self): pass
# ---------- BASIC GUI setup----------
class DataAnalyzerGUI(ctk.CTk):
    def on_closing(self):
        print("Shutting down...")
        os._exit(0)       
    def __init__(self):
        super().__init__()
        self.title("Simple Machine learining (SML) BETA BY VSB-FEI 2026")
        self.geometry("1400x1000")
        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        
        try:
            icon_img = Image.new("RGBA", (256, 256), (0, 188, 212))
            draw = ImageDraw.Draw(icon_img)
            
            # White bars on turquoise background
            white = (255, 255, 255)
            bar_width = 10
            bar_spacing = 45
            start_x = 45
            base_y = 205
            bar_height = 110
            
            # 5 bars: same height, outer two positioned 40% higher, center 40% lower
            y_offsets = [
                -0.40 * bar_height,   # Left outer: 40% higher
                0,                     # Left inner: normal
                0.40 * bar_height,    # Center: 40% lower
                0,                     # Right inner: normal
                -0.40 * bar_height    # Right outer: 40% higher
            ]
            
            for i, y_offset in enumerate(y_offsets):
                x0 = start_x + i * bar_spacing
                y0 = int(base_y - bar_height + y_offset)
                x1 = x0 + bar_width
                y1 = int(y0 + bar_height)
                draw.rounded_rectangle([x0, y0, x1, y1], radius=4, fill=white)
            
            icon_img.save("app_icon.ico")
            self.iconbitmap(default="app_icon.ico")
        except Exception as e:
            print(f"Could not set app icon: {e}")
            
#Variables
        self.df = None
        self.original_path = None
        self.appended_files = []
        self.appended_file_counts = []
        self.appended_file_paths = []
        self.scoring_df = None
        self.input_cols = []
        self.output_col = None
        self.current_selected_col = None
        self.list_widgets = {} 
        self.hyper_widgets = {}
        self.cancel_event = threading.Event()
        self.training_thread = None
        self.last_model_summary = ""
        self.last_results = {}
        self.epoch_history = None
        self.partition_mode_var = ctk.StringVar(value="partition")
        self.external_test_df = None
        self.external_test_path = None
#Layout
        self.grid_columnconfigure(0, weight=3)
        self.grid_columnconfigure(1, weight=0, minsize=500) 
        self.grid_rowconfigure(0, weight=1)
# LEFT: Data Table & Console
        self.left_frame = ctk.CTkFrame(self)
        self.left_frame.grid(row=0, column=0, padx=20, pady=20, sticky="nsew")
        self.left_frame.grid_rowconfigure(0, weight=0)
        self.left_frame.grid_rowconfigure(1, weight=0)
        self.left_frame.grid_rowconfigure(2, weight=1)
        self.left_frame.grid_rowconfigure(3, weight=1)
        self.left_frame.grid_rowconfigure(4, weight=0)
        self.left_frame.grid_columnconfigure(0, weight=1)
        btn_row = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        btn_row.grid(row=0, column=0, padx=10, pady=(10, 5), sticky="ew")
        self.load_btn = ctk.CTkButton(btn_row, text="📁 Load", command=self.load_file, height=40)
        self.load_btn.pack(side="left", expand=True, padx=2, fill="x")
        self.append_btn = ctk.CTkButton(btn_row, text="➕ Append", command=self.append_file, height=40, state="disabled")
        self.append_btn.pack(side="left", expand=True, padx=2, fill="x")
        self.remove_last_btn = ctk.CTkButton(btn_row, text="🧹 Remove Last", command=self.remove_last_appended_file, height=40, state="disabled")
        self.remove_last_btn.pack(side="left", expand=True, padx=2, fill="x")
        self.clear_data_btn = ctk.CTkButton(btn_row, text="❌ Clear All", command=self.clear_all_datasets, height=40, state="disabled")
        self.clear_data_btn.pack(side="left", expand=True, padx=2, fill="x")
        self.appended_files_frame = ctk.CTkScrollableFrame(self.left_frame, height=100, fg_color="#1a1a1a")
        self.appended_files_frame.grid(row=1, column=0, padx=10, pady=(0, 10), sticky="nsew")
        self.appended_files_frame.grid_columnconfigure(0, weight=1)
        self.appended_files_status = ctk.CTkLabel(self.appended_files_frame, text="No datasets loaded", anchor="nw", justify="left")
        self.appended_files_status.pack(fill="both", expand=True, padx=5, pady=5)
        self.tree_container = ctk.CTkFrame(self.left_frame)
        self.tree_container.grid(row=2, column=0, sticky="nsew", padx=10, pady=(5, 5))
        y_scroll = ttk.Scrollbar(self.tree_container, orient="vertical")
        x_scroll = ttk.Scrollbar(self.tree_container, orient="horizontal")
        self.tree = ttk.Treeview(self.tree_container, show="headings", yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        y_scroll.config(command=self.tree.yview)
        x_scroll.config(command=self.tree.xview)
        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        self.tree_container.grid_rowconfigure(0, weight=1)
        self.tree_container.grid_columnconfigure(0, weight=1)
        self.tree.bind("<ButtonRelease-1>", self.on_table_click)
        self.tree.bind("<Configure>", lambda event: self.auto_resize_tree_columns())
        self.console = ctk.CTkTextbox(self.left_frame, wrap="word", font=("Consolas", 12))
        self.console.grid(row=3, column=0, sticky="nsew", padx=10, pady=(5, 10))
        self.after(100, self.start_redirect)

# PHASE 1: SETUP
        self.setup_frame = ctk.CTkFrame(self)
        self.setup_frame.grid(row=0, column=1, padx=(10, 20), pady=20, sticky="nsew")
        ctk.CTkLabel(self.setup_frame, text="Model Setup", font=("Arial", 24, "bold")).pack(pady=15)
        self.sel_info_label = ctk.CTkLabel(self.setup_frame, text="Selected: None", text_color="#3498db")
        self.sel_info_label.pack()
        self.input_scroll_frame = ctk.CTkScrollableFrame(self.setup_frame, height=300, fg_color="#1a1a1a")
        self.input_scroll_frame.pack(padx=15, pady=5, fill="x")
        btn_in_f = ctk.CTkFrame(self.setup_frame, fg_color="transparent")
        btn_in_f.pack(fill="x", padx=15, pady=10)
        ctk.CTkButton(btn_in_f, text="+ Input", fg_color="#2874a6", command=self.add_to_input).pack(side="left", expand=True, padx=5)
        ctk.CTkButton(btn_in_f, text="- Input", fg_color="#922b21", command=self.remove_from_input).pack(side="left", expand=True, padx=5)
        ctk.CTkButton(btn_in_f, text="Data cleaning: Numeric only", fg_color="#8e44ad", command=self.remove_non_numeric_inputs).pack(side="left", expand=True, padx=5)
        self.output_display = ctk.CTkLabel(self.setup_frame, text="Target: None", fg_color="#2c3e50", height=50, corner_radius=8)
        self.output_display.pack(padx=15, pady=5, fill="x")
        ctk.CTkButton(self.setup_frame, text="Set as Target", fg_color="#145a32", command=self.set_output).pack(padx=15, pady=5, fill="x")
        self.next_to_prep_btn = ctk.CTkButton(self.setup_frame, text="NEXT: PREPROCESSING →", height=60, state="disabled", command=self.show_prep)
        self.next_to_prep_btn.pack(side="bottom", pady=(10,5), padx=15, fill="x")
        self.next_to_load_model_btn = ctk.CTkButton(self.setup_frame, text="NEXT: USE MODEL →", height=60, fg_color="#27ae60", command=self.load_saved_model_for_testing)
        self.next_to_load_model_btn.pack(side="bottom", pady=(5,20), padx=15, fill="x")
# PHASE 2: PREPROCESSING
        self.prep_frame = ctk.CTkFrame(self)
        self.prep_scroll = ctk.CTkScrollableFrame(self.prep_frame, fg_color="transparent")
        self.prep_scroll.pack(fill="both", expand=True)
        ctk.CTkLabel(self.prep_scroll, text="Pipeline Settings", font=("Arial", 24, "bold")).pack(pady=15)
        self.extract_var = ctk.BooleanVar(value=False)
        self.extract_mode_var = tk.StringVar(value="sliding")
        self.extract_switch = ctk.CTkSwitch(self.prep_scroll, text="Enable Preprocessing", variable=self.extract_var, command=self.toggle_ext)
        self.extract_switch.pack(pady=10, padx=20, anchor="w")
        self.extra_options_frame = ctk.CTkFrame(self.prep_scroll, fg_color="transparent")
        mode_frame = ctk.CTkFrame(self.extra_options_frame, fg_color="#2b2b2b")
        mode_frame.pack(fill="x", padx=15, pady=(10,5))
        ctk.CTkLabel(mode_frame, text="Extraction Method:", anchor="w").pack(anchor="w", padx=10, pady=(5,0))
        ctk.CTkRadioButton(mode_frame, text="Sliding Window", variable=self.extract_mode_var, value="sliding").pack(anchor="w", padx=20, pady=(5,0))
        ctk.CTkRadioButton(mode_frame, text="Fourier", variable=self.extract_mode_var, value="fourier").pack(anchor="w", padx=20, pady=(5,0))
        ctk.CTkRadioButton(mode_frame, text="Frequency Domain", variable=self.extract_mode_var, value="frequency").pack(anchor="w", padx=20, pady=(5,10))
        win_frame = ctk.CTkFrame(self.extra_options_frame, fg_color="#2b2b2b")
        win_frame.pack(fill="x", padx=15, pady=10)
        ctk.CTkLabel(win_frame, text="Window Size:").grid(row=0, column=0, padx=10, pady=5)
        self.window_size_entry = ctk.CTkEntry(win_frame, width=70)
        self.window_size_entry.insert(0, "100")
        self.window_size_entry.grid(row=0, column=1, padx=10)
        self.window_size_entry.bind("<KeyRelease>", self.update_window_time_info)
        ctk.CTkLabel(win_frame, text="Sampling Freq. (Hz):").grid(row=1, column=0, padx=10, pady=5)
        self.sample_freq_entry = ctk.CTkEntry(win_frame, width=70)
        self.sample_freq_entry.insert(0, "1000")
        self.sample_freq_entry.grid(row=1, column=1, padx=10)
        self.sample_freq_entry.bind("<KeyRelease>", self.update_window_time_info)
        self.window_time_info = ctk.CTkLabel(win_frame, text="Window = 0.10 s (100 samples)", text_color="#aeb6bf")
        self.window_time_info.grid(row=2, column=0, columnspan=2, padx=10, pady=(0, 5), sticky="w")
        ctk.CTkLabel(win_frame, text="Normalize:").grid(row=3, column=0, padx=10, pady=5)
        self.normalize_method_var = tk.StringVar(value="None")
        self.normalize_option_menu = ctk.CTkOptionMenu(win_frame, values=["None", "Standard", "MinMax"], variable=self.normalize_method_var)
        self.normalize_option_menu.grid(row=3, column=1, padx=10, pady=5, sticky="ew")
        ctk.CTkLabel(win_frame, text="Overlap %:").grid(row=4, column=0, padx=10, pady=5)
        self.overlap_slider = ctk.CTkSlider(win_frame, from_=1, to=99, command=self.update_overlap_label)
        self.overlap_slider.set(50)
        self.overlap_slider.grid(row=4, column=1, padx=10, sticky="ew")
        self.overlap_value_lbl = ctk.CTkLabel(win_frame, text="50%")
        self.overlap_value_lbl.grid(row=4, column=2, padx=(0, 10))
        self.update_overlap_label(50)
        self.update_window_time_info()
        self.feats = {
            "Mean": ctk.BooleanVar(value=True),
            "RMS": ctk.BooleanVar(value=True),
            "Std Dev": ctk.BooleanVar(),
            "Min": ctk.BooleanVar(),
            "Max": ctk.BooleanVar(),
            "Peak-to-Peak": ctk.BooleanVar(),
            "Centroid": ctk.BooleanVar(),
            "Bandwidth": ctk.BooleanVar()
        }
        self.feat_checkboxes = []
        feat_grid = ctk.CTkFrame(self.extra_options_frame, fg_color="transparent")
        feat_grid.pack(fill="x", padx=5)
        for i, (name, var) in enumerate(self.feats.items()):
            row, col = divmod(i, 2)
            chk = ctk.CTkCheckBox(feat_grid, text=name, variable=var)
            chk.grid(row=row, column=col, sticky="w", padx=20, pady=5)
            self.feat_checkboxes.append(chk)
        ctk.CTkLabel(self.prep_scroll, text="Test Data Source", font=("Arial", 16, "bold")).pack(pady=(20, 5))
        mode_frame = ctk.CTkFrame(self.prep_scroll, fg_color="transparent")
        mode_frame.pack(fill="x", padx=20)
        self.partition_radio_btn = ctk.CTkRadioButton(mode_frame, text="Partition Dataset", variable=self.partition_mode_var, value="partition", command=self.update_test_source_ui)
        self.partition_radio_btn.pack(side="left", padx=5, pady=5)
        self.external_radio_btn = ctk.CTkRadioButton(mode_frame, text="External Test File", variable=self.partition_mode_var, value="external", command=self.update_test_source_ui)
        self.external_radio_btn.pack(side="left", padx=5, pady=5)
        self.external_test_frame = ctk.CTkFrame(self.prep_scroll, fg_color="transparent")
        self.external_test_frame.pack(fill="x", padx=20, pady=(5,10))
        self.load_external_test_btn = ctk.CTkButton(self.external_test_frame, text="Load External Test File", fg_color="#2874a6", command=self.load_external_test_file)
        self.load_external_test_btn.pack(fill="x")
        self.external_test_status = ctk.CTkLabel(self.external_test_frame, text="No external test file loaded", font=("Consolas", 11), anchor="w")
        self.external_test_status.pack(fill="x", pady=(5,0))
        ctk.CTkLabel(self.prep_scroll, text="Data Partitioning", font=("Arial", 16, "bold")).pack(pady=(10, 5))
        self.split_slider = ctk.CTkSlider(self.prep_scroll, from_=50, to=90, command=self.update_split); self.split_slider.set(70); self.split_slider.pack(padx=20, fill="x")
        self.split_lbl = ctk.CTkLabel(self.prep_scroll, text="70% / 15% / 15%"); self.split_lbl.pack()
        self.update_test_source_ui()
        self.export_feat_btn = ctk.CTkButton(self.prep_scroll, text="Save Extracted Features", fg_color="#2874a6", command=self.instant_export_features)
        self.export_feat_btn.pack(padx=20, pady=10, fill="x")
        self.next_to_model_btn = ctk.CTkButton(self.prep_frame, text="NEXT: ARCHITECT →", height=60, command=self.show_model)
        self.next_to_model_btn.pack(side="bottom", pady=(10,5), padx=15, fill="x")
        ctk.CTkButton(self.prep_frame, text="← Back", command=self.show_setup).pack(side="bottom", padx=15, fill="x")
# PHASE 3: ARCHITECT
        self.model_frame = ctk.CTkFrame(self)
        self.model_scroll = ctk.CTkScrollableFrame(self.model_frame, fg_color="transparent")
        self.model_scroll.pack(fill="both", expand=True)
        ctk.CTkLabel(self.model_scroll, text="Model Configuration", font=("Arial", 24, "bold")).pack(pady=15)
        self.target_type_var = ctk.StringVar(value="Regression")
        self.target_type_selector = ctk.CTkSegmentedButton(self.model_scroll, values=["Regression", "Binary", "Multi-Class"], variable=self.target_type_var)
        self.target_type_selector.pack(pady=(10, 5), padx=20, fill="x")
        self.target_type_hint = ctk.CTkLabel(
            self.model_scroll,
            text="Task options will update based on the selected model.",
            font=("Consolas", 11),
            anchor="w"
        )
        self.target_type_hint.pack(padx=20, pady=(0, 5), fill="x")
        self.model_choice = ctk.StringVar(value="ANN (MLP)")
        self.model_option_menu = ctk.CTkOptionMenu(self.model_scroll, values=list(MODEL_LIBRARY.keys()), variable=self.model_choice, command=self.build_hyper_ui)
        self.model_option_menu.pack(pady=10, padx=20, fill="x")
        self.hyper_container = ctk.CTkFrame(self.model_scroll, fg_color="#2b2b2b")
        self.hyper_container.pack(fill="x", padx=15, pady=15)
        self.build_hyper_ui("ANN (MLP)") 
        self.status_label = ctk.CTkLabel(self.model_frame, text="Ready to train", font=("Consolas", 12), anchor="w")
        self.status_label.pack(padx=15, pady=(5, 10), fill="x")
        self.progress_label = ctk.CTkLabel(self.model_frame, text="Progress: idle", font=("Consolas", 11), anchor="w")
        self.progress_label.pack(padx=15, pady=(0, 3), fill="x")
        self.progress_var = tk.DoubleVar(value=0.0)
        self.training_progress = ctk.CTkProgressBar(self.model_frame, orientation="horizontal", mode="determinate", variable=self.progress_var)
        self.training_progress.pack(padx=15, pady=(0, 10), fill="x")
        self.run_btn = ctk.CTkButton(self.model_frame, text="🚀 START TRAINING", height=60, fg_color="#d35400", command=self.execute)
        self.run_btn.pack(side="bottom", pady=(10, 8), padx=15, fill="x")
        self.grid_search_btn = ctk.CTkButton(self.model_frame, text="🔎 GRID SEARCH", height=40, fg_color="#2980b9", command=self.perform_grid_search)
        self.grid_search_btn.pack(side="bottom", pady=(0, 8), padx=15, fill="x")
        self.cancel_btn = ctk.CTkButton(self.model_frame, text="🛑 STOP TRAINING", height=40, fg_color="#c0392b", state="disabled", command=self.cancel_training)
        self.cancel_btn.pack(side="bottom", pady=(0, 10), padx=15, fill="x")
        ctk.CTkButton(self.model_frame, text="← Back", command=self.show_prep).pack(side="bottom", padx=15, fill="x")
# PHASE 4: POST-TRAINING
        self.post_frame = ctk.CTkFrame(self)
        ctk.CTkLabel(self.post_frame, text="Results & Export", font=("Arial", 24, "bold")).pack(pady=15)
        self.metrics_label = ctk.CTkLabel(self.post_frame, text="Model Not Trained", font=("Consolas", 14), justify="left")
        self.metrics_label.pack(pady=10, padx=20)
        btn_f = ctk.CTkFrame(self.post_frame, fg_color="transparent")
        btn_f.pack(fill="x", padx=15, pady=10)
        ctk.CTkButton(btn_f, text="📈 Plot Results", command=self.plot_results).pack(pady=5, fill="x")
        ctk.CTkButton(btn_f, text="� Accuracy / Epoch", fg_color="#3498db", command=self.plot_accuracy_per_epoch).pack(pady=5, fill="x")
        ctk.CTkButton(btn_f, text="�💾 Export Model", fg_color="#27ae60", command=self.save_model).pack(pady=5, fill="x")
        ctk.CTkButton(btn_f, text="📝 Export Results to Word", fg_color="#8e44ad", command=self.export_results_to_word).pack(pady=5, fill="x")
        ctk.CTkButton(self.post_frame, text="← Back to Model", 
              command=self.back_to_tweak).pack(side="bottom", pady=10, padx=15, fill="x")         
        # Button 2: Full Reset
        ctk.CTkButton(self.post_frame, text="🔄 New Dataset / Full Reset", 
              fg_color="#34495e", command=self.show_setup).pack(side="bottom", pady=5, padx=15, fill="x")
 #PHASE 4: POST-TRAINING
    def update_split(self, value):
        train = int(value)
        rem = (100 - train) // 2
        test = 100 - train - rem
        self.split_lbl.configure(text=f"{train}% / {rem}% / {test}%")

    def update_overlap_label(self, value):
        self.overlap_value_lbl.configure(text=f"{int(float(value))}%")

    def update_window_time_info(self, event=None):
        try:
            window_size = int(self.window_size_entry.get())
            sample_freq = float(self.sample_freq_entry.get())
            if sample_freq <= 0:
                raise ValueError
            duration_sec = window_size / sample_freq
            if duration_sec >= 1:
                duration_text = f"{duration_sec:.2f} s"
            else:
                duration_text = f"{duration_sec*1000:.1f} ms"
            self.window_time_info.configure(
                text=f"Window = {duration_text} ({window_size} samples at {sample_freq:.0f} Hz)"
            )
        except Exception:
            self.window_time_info.configure(
                text="Window = invalid size/frequency"
            )

    def apply_normalization_to_sets(self, Xtr, Xv=None, Xt=None):
        method = getattr(self, 'normalize_method_var', tk.StringVar(value='None')).get()
        if method == "None" or Xtr is None:
            return Xtr, Xv, Xt
        scaler = StandardScaler() if method == "Standard" else MinMaxScaler()
        if Xtr.ndim == 3:
            original_shape = Xtr.shape
            Xtr = scaler.fit_transform(Xtr.reshape(-1, original_shape[2])).reshape(original_shape)
        else:
            Xtr = scaler.fit_transform(Xtr)
        if Xv is not None:
            Xv = scaler.transform(Xv.reshape(-1, Xv.shape[2])).reshape(Xv.shape) if Xv.ndim == 3 else scaler.transform(Xv)
        if Xt is not None:
            Xt = scaler.transform(Xt.reshape(-1, Xt.shape[2])).reshape(Xt.shape) if Xt.ndim == 3 else scaler.transform(Xt)
        return Xtr, Xv, Xt

    def _read_tabular_file(self, file_path, header='infer'):
        if file_path.endswith('.csv'):
            candidates = [',', ';', '\t', '|']
            try:
                with open(file_path, newline='', encoding='utf-8', errors='ignore') as f:
                    sample = f.read(4096)
                    if sample.strip():
                        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
                        candidates = [dialect.delimiter] + [c for c in [',', ';', '\t', '|'] if c != dialect.delimiter]
            except Exception:
                pass
            last_error = None
            for sep in candidates:
                for decimal in ('.', ','):
                    if sep == decimal:
                        continue
                    try:
                        df = pd.read_csv(
                            file_path,
                            sep=sep,
                            engine='python',
                            header=header,
                            decimal=decimal,
                            thousands='.' if decimal == ',' else None
                        )
                        df = df.copy()
                        for col in df.columns:
                            df[col] = normalize_numeric_series(df[col])
                        return df
                    except Exception as e:
                        last_error = e
            if last_error is not None:
                raise last_error
            raise ValueError(f"Could not parse CSV file: {file_path}")
        else:
            df = pd.read_excel(file_path, header=header)
            for col in df.columns:
                df[col] = normalize_numeric_series(df[col])
            return df
    def append_file(self):
        file_paths = filedialog.askopenfilenames(
            filetypes=[("CSV Files", "*.csv"), ("Excel Files", "*.xlsx *.xls")]
        )
        if not file_paths:
            return
        if self.df is None:
            messagebox.showerror("Append Error", "Load a base dataset first before appending additional files.")
            return
        try:
            existing_paths = {os.path.abspath(self.original_path)} | set(self.appended_file_paths)
            added_files = []
            skipped_files = []
            for file_path in file_paths:
                abs_file_path = os.path.abspath(file_path)
                if abs_file_path in existing_paths:
                    skipped_files.append(os.path.basename(file_path))
                    continue
                existing_paths.add(abs_file_path)
                candidate_df = self._read_tabular_file(file_path, header=0)
                if candidate_df.shape[1] != self.df.shape[1] or self._is_default_header(candidate_df.columns):
                    candidate_df = self._read_tabular_file(file_path, header=None)
                if candidate_df.shape[1] != self.df.shape[1]:
                    raise ValueError(f"The appended file '{os.path.basename(file_path)}' must have the same number of columns as the base file.")
                if not candidate_df.empty and self._looks_like_header_row(candidate_df.iloc[0], self.df.columns):
                    candidate_df = candidate_df.iloc[1:].reset_index(drop=True)
                if candidate_df.empty:
                    raise ValueError(f"Appended file '{os.path.basename(file_path)}' contains no rows after removing the header row.")
                candidate_df.columns = self.df.columns
                self.df = pd.concat([self.df, candidate_df], ignore_index=True)
                self.appended_files.append(os.path.basename(file_path))
                self.appended_file_paths.append(abs_file_path)
                self.appended_file_counts.append(len(candidate_df))
                added_files.append(os.path.basename(file_path))
            self.input_cols = list(self.df.columns)
            self._update_appended_files_ui()
            self._update_data_buttons()
            self._refresh_data_table()
            self.external_test_df = None
            self.external_test_path = None
            self.external_test_status.configure(text="No external test file loaded")
            self.output_display.configure(text="Target: None")
            self.refresh_setup()
            if added_files:
                print(f">> Appended: {', '.join(added_files)}.")
            if skipped_files:
                messagebox.showwarning(
                    "Duplicate files skipped",
                    "The following file(s) were already loaded or appended and were skipped:\n" + "\n".join(skipped_files)
                )
        except Exception as e:
            messagebox.showerror("Append Error", f"Could not append file: {e}")

    def _looks_like_header_row(self, row, columns):
        try:
            values = [str(x).strip().lower() for x in row.tolist()]
            col_names = [str(c).strip().lower() for c in columns]
            if len(values) != len(col_names):
                return False
            match_count = sum(1 for value in values if value in col_names)
            if match_count >= max(2, len(col_names) // 2):
                return True
            numeric_count = sum(1 for value in values if self._is_numeric_string(value))
            if numeric_count == 0 and any(value for value in values):
                return True
        except Exception:
            return False
        return False

    def _is_numeric_string(self, value):
        if value is None:
            return False
        try:
            float(str(value).replace(',', '.'))
            return True
        except Exception:
            return False

    def _is_default_header(self, columns):
        normalized = [str(c).strip().lower() for c in columns]
        if all(c.isdigit() for c in normalized):
            return True
        if all(c.startswith("unnamed") for c in normalized):
            return True
        return False

    def _update_appended_files_ui(self, added_file=None):
        # Optionally add a newly appended file
        if added_file:
            self.appended_files.append(added_file)
        # Build list beginning with the base/loaded file (if any)
        files = []
        if getattr(self, 'original_path', None):
            files.append(os.path.basename(self.original_path))
        files.extend(self.appended_files)
        if files:
            display_text = "Loaded files:\n" + "\n".join(files)
        else:
            display_text = "No datasets loaded"
        self.appended_files_status.configure(text=display_text)

    def _update_data_buttons(self):
        has_main = self.df is not None
        has_appends = bool(self.appended_files)
        if hasattr(self, 'append_btn'):
            self.append_btn.configure(state="normal" if has_main else "disabled")
        if hasattr(self, 'clear_data_btn'):
            self.clear_data_btn.configure(state="normal" if has_main else "disabled")
        if hasattr(self, 'remove_last_btn'):
            self.remove_last_btn.configure(state="normal" if has_appends else "disabled")

    def remove_last_appended_file(self):
        if not self.appended_files or not self.appended_file_counts:
            return
        last_count = self.appended_file_counts.pop()
        removed_file = self.appended_files.pop()
        if last_count > 0 and len(self.df) >= last_count:
            self.df = self.df.iloc[:-last_count].reset_index(drop=True)
        self._update_appended_files_ui()
        self._update_data_buttons()
        self._refresh_data_table()
        print(f">> Removed last appended file: {removed_file}")

    def clear_all_datasets(self):
        self.df = None
        self.original_path = None
        self.appended_files = []
        self.appended_file_counts = []
        self.appended_file_paths = []
        self.input_cols = []
        self.output_col = None
        self.current_selected_col = None
        self.tree["columns"] = []
        self.tree.delete(*self.tree.get_children())
        self._update_appended_files_ui()
        self._update_data_buttons()
        self.external_test_df = None
        self.external_test_path = None
        self.external_test_status.configure(text="No external test file loaded")
        self.output_display.configure(text="Target: None")
        self.refresh_setup()
        print(">> Cleared all loaded datasets.")

    def _refresh_data_table(self):
        self.tree["columns"] = list(self.df.columns)
        self.tree["show"] = "headings"
        self.tree.delete(*self.tree.get_children())
        for column in self.tree["columns"]:
            self.tree.heading(column, text=column)
            self.tree.column(column, width=120, anchor="center", stretch=False)
        for row in self.df.head(100).to_numpy().tolist():
            self.tree.insert("", "end", values=row)
        self.auto_resize_tree_columns()

    def load_external_test_file(self):
        file_path = filedialog.askopenfilename(
            filetypes=[("CSV Files", "*.csv"), ("Excel Files", "*.xlsx *.xls")]
        )
        if not file_path:
            return
        try:
            external_df = self._read_tabular_file(file_path)
            if self.output_col is not None and self.output_col not in external_df.columns:
                raise ValueError("External test file does not contain the selected target column.")
            for col in self.input_cols:
                if col not in external_df.columns:
                    raise ValueError("External test file does not contain one or more selected input columns.")
            self.external_test_df = external_df
            self.external_test_path = file_path
            self.external_test_status.configure(text=f"Loaded: {os.path.basename(file_path)} ({len(external_df)} rows)")
            print(f">> Loaded external test file: {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Load Error", f"Could not read external test file: {e}")
    def update_test_source_ui(self):
        if self.partition_mode_var.get() == "external":
            self.external_test_frame.pack(fill="x", padx=20, pady=(5,10))
            self.split_slider.configure(state="disabled")
            self.split_lbl.configure(text="External test dataset selected")
        else:
            self.external_test_frame.pack_forget()
            self.split_slider.configure(state="normal")
            self.update_split(self.split_slider.get())
    def toggle_ext(self):
        if getattr(self, 'extract_var', None) and self.extract_var.get():
            try:
                self.extra_options_frame.pack(after=self.extract_switch, fill="x", padx=15, pady=5)
            except Exception:
                self.extra_options_frame.pack(fill="x", padx=15, pady=5)
        else:
            try:
                self.extra_options_frame.pack_forget()
            except Exception:
                pass

    def _get_allowed_target_types(self, model_name):
        if model_name == "LR (Linear/Logistic)":
            return ["Regression", "Binary"]
        return ["Regression", "Binary", "Multi-Class"]

    def update_target_type_constraints(self, model_name=None):
        if model_name is None:
            model_name = self.model_choice.get()
        allowed = self._get_allowed_target_types(model_name)
        if hasattr(self, 'target_type_selector') and hasattr(self, 'target_type_var'):
            self.target_type_selector.configure(values=allowed)
            current = self.target_type_var.get()
            if current not in allowed:
                self.target_type_var.set(allowed[0])
            if hasattr(self, 'target_type_hint'):
                self.target_type_hint.configure(
                    text=f"Task options for {model_name}: {', '.join(allowed)}"
                )

    def build_hyper_ui(self, model_name):
        for w in getattr(self, 'hyper_container', []).winfo_children():
            w.destroy()
        self.hyper_widgets = {}
        params = MODEL_LIBRARY.get(model_name, {})
        for name, info in params.items():
            row = ctk.CTkFrame(self.hyper_container, fg_color="transparent")
            row.pack(fill="x", padx=10, pady=5)
            ctk.CTkLabel(row, text=name, width=150, anchor="w").pack(side="left")
            if info.get("type") == "list":
                w = ctk.CTkOptionMenu(row, values=info.get("options", []))
                if info.get("options"):
                    try:
                        w.set(info.get("options")[0])
                    except Exception:
                        pass
            else:
                w = ctk.CTkEntry(row)
                default = info.get("default", "")
                try:
                    w.insert(0, str(default))
                except Exception:
                    pass
            w.pack(side="right", expand=True, fill="x")
            self.hyper_widgets[name] = w
        self.update_target_type_constraints(model_name)
    def start_redirect(self):
        """Delayed redirection to prevent Tkinter attribute errors."""
        self.redirector = ConsoleRedirect(self.console)
        sys.stdout = self.redirector
        sys.stderr = self.redirector
        print(">> System: console redirected.")      
    def load_file(self):
        file_paths = filedialog.askopenfilenames(
            filetypes=[("CSV Files", "*.csv"), ("Excel Files", "*.xlsx *.xls")]
        )
        if not file_paths:
            return
        try:
            # Load first file as base
            first = file_paths[0]
            base_df = self._read_tabular_file(first, header=0)
            # If first file looks like it used a default header (numbers/Unnamed), try header=None
            if self._is_default_header(base_df.columns):
                base_df = self._read_tabular_file(first, header=None)
                # if header row detected in first row, set columns and drop that row
                if self._looks_like_header_row(base_df.iloc[0], base_df.columns):
                    header = [str(x) for x in base_df.iloc[0].tolist()]
                    base_df = base_df.iloc[1:].reset_index(drop=True)
                    base_df.columns = header

            self.appended_files = []
            self.appended_file_counts = []
            self.appended_file_paths = []
            seen_paths = {os.path.abspath(first)}
            skipped_files = []
            # Iterate remaining files and append after validation
            for fp in file_paths[1:]:
                abs_fp = os.path.abspath(fp)
                if abs_fp in seen_paths:
                    skipped_files.append(os.path.basename(fp))
                    continue
                seen_paths.add(abs_fp)
                cand = self._read_tabular_file(fp, header=0)
                if cand.shape[1] != base_df.shape[1] or self._is_default_header(cand.columns):
                    cand = self._read_tabular_file(fp, header=None)
                # If the first row of candidate looks like a header, drop it
                if not cand.empty and self._looks_like_header_row(cand.iloc[0], base_df.columns):
                    cand = cand.iloc[1:].reset_index(drop=True)
                # If still mismatched number of columns, raise
                if cand.shape[1] != base_df.shape[1]:
                    raise ValueError(f"File '{os.path.basename(fp)}' has a different number of columns.")
                # Force candidate columns to match base
                cand.columns = base_df.columns
                base_df = pd.concat([base_df, cand], ignore_index=True)
                self.appended_files.append(os.path.basename(fp))
                self.appended_file_paths.append(abs_fp)
                self.appended_file_counts.append(len(cand))
            if skipped_files:
                messagebox.showwarning(
                    "Duplicate files ignored",
                    "The following file(s) were already selected and were ignored:\n" + "\n".join(skipped_files)
                )

            self.df = base_df
            self.original_path = file_paths[0]
            self.tree["columns"] = list(self.df.columns)
            self.tree["show"] = "headings"
            self.tree.delete(*self.tree.get_children())
            for column in self.tree["columns"]:
                self.tree.heading(column, text=column)
                self.tree.column(column, width=120, anchor="center", stretch=False)

            # Fill table with first 100 rows for performance
            df_rows = self.df.head(100).to_numpy().tolist()
            for row in df_rows:
                self.tree.insert("", "end", values=row)
            self.auto_resize_tree_columns()
            print(f">> Loaded: {os.path.basename(self.original_path)} with {len(self.df)} rows.")
            self.input_cols = list(self.df.columns)
            print(">> ALL variables selected as input. Select target and remove unwanted inputs.")
            self.output_col = None
            self.external_test_df = None
            self.external_test_path = None
            self._update_appended_files_ui()
            self._update_data_buttons()
            self.external_test_status.configure(text="No external test file loaded")
            self.output_display.configure(text="Target: None")
            self.refresh_setup()
        except Exception as e:
            messagebox.showerror("Load Error", f"Could not read file(s): {e}")
        if not self.tree_container.winfo_ismapped():
            self.tree_container.grid(row=1, column=0, sticky="nsew", padx=10, pady=(10, 5))
        if not self.console.winfo_ismapped():
            self.console.grid(row=2, column=0, sticky="nsew", padx=10, pady=(5, 10))
        if hasattr(self, 'run_btn'):
            self.run_btn.configure(state="normal", text="🚀 START TRAINING", fg_color="#d35400")
            if hasattr(self, 'status_label'):
                self.status_label.configure(text="Ready to train")
    def set_training_state(self, running: bool):
        if not hasattr(self, 'run_btn'):
            return
        control_widgets = [
            getattr(self, 'model_option_menu', None),
            getattr(self, 'target_type_selector', None),
            getattr(self, 'split_slider', None),
            getattr(self, 'extract_switch', None),
            getattr(self, 'window_size_entry', None),
            getattr(self, 'overlap_slider', None),
            getattr(self, 'partition_radio_btn', None),
            getattr(self, 'external_radio_btn', None),
            getattr(self, 'load_external_test_btn', None),
            getattr(self, 'export_feat_btn', None),
            getattr(self, 'next_to_model_btn', None),
            getattr(self, 'next_to_prep_btn', None),
            getattr(self, 'grid_search_btn', None)
        ]
        state_value = "disabled" if running else "normal"
        for widget in control_widgets:
            if widget is not None and hasattr(widget, 'configure'):
                widget.configure(state=state_value)
        for widget in self.hyper_widgets.values():
            if hasattr(widget, 'configure'):
                widget.configure(state=state_value)
        for chk in getattr(self, 'feat_checkboxes', []):
            if hasattr(chk, 'configure'):
                chk.configure(state=state_value)
        if running:
            self.run_btn.configure(state="disabled", text="⏳ Training...", fg_color="#7f8c8d")
            if hasattr(self, 'cancel_btn'):
                self.cancel_btn.configure(state="normal")
            if hasattr(self, 'status_label'):
                self.status_label.configure(text="Training in progress. Please wait...")
            if hasattr(self, 'progress_var'):
                self.progress_var.set(0.0)
            if hasattr(self, 'progress_label'):
                self.progress_label.configure(text="Starting...")
        else:
            self.run_btn.configure(state="normal", text="🚀 START TRAINING", fg_color="#d35400")
            if hasattr(self, 'cancel_btn'):
                self.cancel_btn.configure(state="disabled")
            if hasattr(self, 'progress_var'):
                self.progress_var.set(1.0)
            if hasattr(self, 'cancel_event') and self.cancel_event.is_set():
                if hasattr(self, 'status_label'):
                    self.status_label.configure(text="Training canceled. Check results below.")
                if hasattr(self, 'progress_label'):
                    self.progress_label.configure(text="Canceled")
            else:
                if hasattr(self, 'status_label'):
                    self.status_label.configure(text="Training complete. Check results below.")
                if hasattr(self, 'progress_label'):
                    self.progress_label.configure(text="Complete")
    def update_progress(self, current, total=1, text=None):
        if not hasattr(self, 'training_progress'):
            return
        try:
            ratio = float(current) / float(total) if total else 0.0
        except Exception:
            ratio = 0.0
        ratio = max(0.0, min(1.0, ratio))
        if hasattr(self, 'progress_var'):
            self.progress_var.set(ratio)
        if hasattr(self, 'training_progress'):
            self.training_progress.update_idletasks()
        if hasattr(self, 'status_label') and text is not None:
            self.status_label.configure(text=text)
        if hasattr(self, 'progress_label') and text is not None:
            self.progress_label.configure(text=text)

    def cancel_training(self):
        if hasattr(self, 'cancel_event'):
            self.cancel_event.set()
            print(">> Training cancellation requested.")
            if hasattr(self, 'status_label'):
                self.status_label.configure(text="Cancel requested... waiting for training loop to stop.")
    def show_setup(self):
        self.reset_ui_layout()
        self.prep_frame.grid_forget()
        self.model_frame.grid_forget()
        self.post_frame.grid_forget()
        self.setup_frame.grid(row=0, column=1, padx=20, pady=20, sticky="nsew")
    def show_prep(self):
        self.reset_ui_layout()
        self.setup_frame.grid_forget()
        self.model_frame.grid_forget()
        self.post_frame.grid_forget()
        self.prep_frame.grid(row=0, column=1, padx=20, pady=20, sticky="nsew")

    def show_post_training(self):
        try:
            self.model_frame.grid_forget()
        except Exception:
            pass
        try:
            self.post_frame.grid(row=0, column=1, padx=20, pady=20, sticky="nsew")
        except Exception:
            pass

    def reset_ui_layout(self):
        if not self.tree_container.winfo_ismapped():
            self.tree_container.grid(row=1, column=0, sticky="nsew", padx=10, pady=(10, 5))
        if not self.console.winfo_ismapped():
            self.console.grid(row=2, column=0, sticky="nsew", padx=10, pady=(5, 10))
        if hasattr(self, 'run_btn'):
            try:
                self.run_btn.configure(state="normal", text="🚀 START TRAINING", fg_color="#d35400")
            except Exception:
                pass
            if hasattr(self, 'status_label'):
                try:
                    self.status_label.configure(text="Ready to train")
                except Exception:
                    pass

    def show_model(self):
        self.reset_ui_layout()
        self.prep_frame.grid_forget()
        self.post_frame.grid_forget()
        self.model_frame.grid(row=0, column=1, padx=20, pady=20, sticky="nsew")
    def back_to_tweak(self):
        if hasattr(self, 'metrics_label'):
            try:
                self.metrics_label.configure(text="Ready for re-training...")
            except Exception:
                pass
        # Clear trained model reference if present
        if hasattr(self, 'trained_model'):
            self.trained_model = None
        self.show_model()

    def refresh_setup(self):
        for w in self.input_scroll_frame.winfo_children():
            w.destroy()
        if self.df is None:
            return
        for c in self.df.columns:
            is_selected = c == self.current_selected_col
            is_input = c in self.input_cols
            is_target = c == self.output_col
            if is_target:
                text = f"🎯 {c}"
                fg_color = "#1e8449"
            elif is_selected:
                text = f"▶ {c}"
                fg_color = "#2e86de"
            elif is_input:
                text = f"📥 {c}"
                fg_color = "#5d6d7e"
            else:
                text = f"○ {c}"
                fg_color = "#34495e"
            btn = ctk.CTkButton(
                self.input_scroll_frame,
                text=text,
                fg_color=fg_color,
                anchor="w",
                command=lambda x=c: self.set_active(x)
            )
            btn.pack(fill="x", pady=1)
        if hasattr(self, 'next_to_prep_btn'):
            self.next_to_prep_btn.configure(state="normal" if (self.input_cols and self.output_col) else "disabled")
    def set_active(self, c): 
        self.current_selected_col = c
        self.sel_info_label.configure(text=f"Selected: {c}")
        self.refresh_setup()
    def add_to_input(self): 
        if self.current_selected_col and self.current_selected_col not in self.input_cols:
            self.input_cols.append(self.current_selected_col)
            self.refresh_setup()
    def remove_from_input(self):
        if self.current_selected_col in self.input_cols:
            self.input_cols.remove(self.current_selected_col); self.refresh_setup()
    def remove_non_numeric_inputs(self):
        if self.df is None:
            messagebox.showerror("Error", "Load a dataset first before cleaning inputs.")
            return
        cleaned = [col for col in self.input_cols if pd.api.types.is_numeric_dtype(self.df[col])]
        removed = [col for col in self.input_cols if col not in cleaned]
        self.input_cols = cleaned
        self.refresh_setup()
        if removed:
            messagebox.showinfo("Data Cleaning", f"Removed non-numeric inputs: {', '.join(removed)}")
        else:
            messagebox.showinfo("Data Cleaning", "No non-numeric input columns were found.")
    def set_output(self):
        if self.current_selected_col:
            self.output_col = self.current_selected_col
            if self.output_col in self.input_cols: self.input_cols.remove(self.output_col)
            self.output_display.configure(text=f"Target: {self.output_col}"); self.refresh_setup()
    def on_table_click(self, event):
        col = self.tree.identify_column(event.x)
        if col: self.set_active(self.df.columns[int(col[1:])-1])
    def auto_resize_tree_columns(self):
        try:
            tree_font = font.Font(font=self.tree.cget("font"))
        except Exception:
            tree_font = font.nametofont("TkDefaultFont")
        for col in self.tree["columns"]:
            max_width = tree_font.measure(col) + 20
            for item in self.tree.get_children():
                cell_text = str(self.tree.set(item, col))
                max_width = max(max_width, tree_font.measure(cell_text) + 20)
            self.tree.column(col, width=max_width, stretch=False)
    def update_post_training_ui(self, y_true, y_pred):
        self.last_y_true, self.last_y_pred = y_true, y_pred
        if self.target_type_var.get() == "Regression":
            mse = mean_squared_error(y_true, y_pred)
            r2 = r2_score(y_true, y_pred)
            stats = f"MSE: {mse:.4f}\nR2 Score: {r2:.4f}"
        else:
            acc = (y_true == y_pred).mean()
            stats = f"Accuracy: {acc*100:.2f}%"
        self.metrics_label.configure(text=stats)
        self.show_post_training()
    def instant_export_features(self):
        if self.df is None:
            messagebox.showerror("Error", "Load a dataset first.")
            return
        if self.output_col is None or not self.input_cols:
            messagebox.showerror("Error", "Select input and target columns before exporting features.")
            return
        save_path = filedialog.asksaveasfilename(defaultextension=".csv")
        if not save_path:
            return
        try:
            feats_config = {k: v.get() for k, v in self.feats.items()}
            X, y = sliding_window_features(
                self.df,
                self.input_cols,
                self.output_col,
                int(self.window_size_entry.get()),
                int(self.overlap_slider.get()),
                feats_config
            )
            selected_features = [name for name, enabled in feats_config.items() if enabled]
            feature_columns = []
            for col in self.input_cols:
                for feat_name in selected_features:
                    feature_columns.append(f"{col}_{feat_name}")
            if self.normalize_method_var.get() != "None":
                scaler = StandardScaler() if self.normalize_method_var.get() == "Standard" else MinMaxScaler()
                X = scaler.fit_transform(X)
                export_df = pd.DataFrame(X, columns=feature_columns)
            else:
                export_df = pd.DataFrame(X, columns=feature_columns)
            export_df[self.output_col] = y
            export_df.to_csv(save_path, index=False)
            messagebox.showinfo("Success", "Exported!")
        except Exception as e:
            messagebox.showerror("Error", str(e))
    def save_model(self):
        model_type = self.model_choice.get()
        ext = ".pth" if model_type == "ANN (MLP)" else ".pkl"
        path = filedialog.asksaveasfilename(defaultextension=ext)
        if path:
            if model_type == "ANN (MLP)":
                torch.save({
                    'state_dict': self.trained_model.state_dict(),
                    'input_dim': self.trained_model.net[0].in_features if hasattr(self.trained_model, 'net') else None,
                    'num_outputs': self.trained_model.net[-1].out_features if hasattr(self.trained_model, 'net') else None,
                    'hidden_layers': int(self.hyper_widgets.get('Hidden Layers').get()) if 'Hidden Layers' in self.hyper_widgets else None,
                    'neurons': int(self.hyper_widgets.get('Neurons').get()) if 'Neurons' in self.hyper_widgets else None,
                    'activation': self.hyper_widgets.get('Activation').get() if 'Activation' in self.hyper_widgets else 'ReLU'
                }, path)
            else:
                with open(path, 'wb') as f: pickle.dump(self.trained_model, f)
            print(f"Model saved to {path}")
    def load_saved_model_for_testing(self):
        if self.df is None:
            messagebox.showerror("Error", "No dataset loaded. Please load a CSV or Excel file first.")
            return
        file_path = filedialog.askopenfilename(filetypes=[("Model files", "*.pth *.pkl"), ("PyTorch models", "*.pth"), ("Pickle models", "*.pkl")])
        if not file_path:
            return
        try:
            target_type = self.target_type_var.get()
            test_df = self.external_test_df if self.partition_mode_var.get() == "external" and self.external_test_df is not None else self.df
            if self.input_cols:
                X = test_df[self.input_cols].values
            else:
                if self.output_col and self.output_col in test_df.columns:
                    X = test_df.drop(columns=[self.output_col]).values
                else:
                    X = test_df.values
            y = None
            if self.output_col and self.output_col in test_df.columns:
                y = test_df[self.output_col].values
            effective_target_type = target_type
            loaded_target_type_info = None
            if file_path.lower().endswith('.pth'):
                if self.model_choice.get() != "ANN (MLP)":
                    raise ValueError("Loaded .pth files require ANN (MLP) model selection.")
                saved_num_outputs = len(np.unique(y)) if (y is not None and target_type != "Regression") else None
                model, saved_num_outputs = self._load_saved_ann_model(file_path, X.shape[1], target_type, saved_num_outputs)
                if saved_num_outputs is not None and saved_num_outputs > 1 and target_type == "Regression":
                    effective_target_type = "Binary" if saved_num_outputs == 2 else "Multi-Class"
                    self.target_type_var.set(effective_target_type)
                    loaded_target_type_info = effective_target_type
            else:
                with open(file_path, 'rb') as f:
                    model = pickle.load(f)
            y_pred = self._predict_with_model(model, X, effective_target_type)
            if y_pred is None:
                raise ValueError("Failed to obtain predictions from the loaded model.")
            self.last_model_summary = f"Loaded saved model from: {os.path.basename(file_path)}"
            self.last_results = {}
            self.last_y_true = y
            self.last_y_pred = y_pred
            if y is not None:
                self.print_detailed_metrics(y, y_pred, "TESTING", effective_target_type)
                message = "Saved model evaluated on selected dataset."
                if loaded_target_type_info:
                    message += f" Target type updated to {loaded_target_type_info} based on loaded ANN output dimension."
                    messagebox.showinfo("Target Type Updated", f"Loaded ANN model inferred classification task. Target type set to {loaded_target_type_info}.")
                self.metrics_label.configure(text=message)
            else:
                self.metrics_label.configure(text="Saved model loaded and predictions generated (no target selected).")
            self.show_post_training()
        except Exception as e:
            messagebox.showerror("Load Error", str(e))
    def _prepare_test_data(self, df, target_type):
        feats_config = {k: v.get() for k, v in self.feats.items()}
        if self.extract_var.get():
            mode = self.extract_mode_var.get()
            if mode == "sliding":
                X, y = sliding_window_features(df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats_config)
            elif mode == "fourier":
                X, y = fourier_window_features(df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats_config)
            else:
                X, y = frequency_domain_features(df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats_config)
        else:
            X = df[self.input_cols].values
            y = df[self.output_col].values
        if target_type != "Regression":
            _, y = np.unique(y, return_inverse=True)
        return X, y, None
    def _load_saved_ann_model(self, path, input_dim, target_type, num_outputs=None):
        checkpoint = torch.load(path, map_location="cpu", weights_only=True)
        if isinstance(checkpoint, dict) and 'state_dict' in checkpoint:
            state_dict = checkpoint['state_dict']
            saved_input_dim = checkpoint.get('input_dim', input_dim)
            saved_num_outputs = checkpoint.get('num_outputs', num_outputs)
            saved_hidden_layers = checkpoint.get('hidden_layers', None)
            saved_neurons = checkpoint.get('neurons', None)
            saved_activation = checkpoint.get('activation', None)
        else:
            state_dict = checkpoint
            saved_input_dim = input_dim
            saved_num_outputs = num_outputs
            saved_hidden_layers = None
            saved_neurons = None
            saved_activation = None
            weight_keys = [k for k in state_dict.keys() if k.endswith('.weight')]
            if weight_keys:
                weight_keys.sort()
                first_w = state_dict[weight_keys[0]]
                last_w = state_dict[weight_keys[-1]]
                if first_w.ndim == 2:
                    saved_input_dim = first_w.shape[1]
                if last_w.ndim == 2:
                    saved_num_outputs = last_w.shape[0]
        hidden_layers = int(saved_hidden_layers) if saved_hidden_layers is not None else (int(self.hyper_widgets['Hidden Layers'].get()) if 'Hidden Layers' in self.hyper_widgets else int(MODEL_LIBRARY['ANN (MLP)']['Hidden Layers']['default']))
        neurons = int(saved_neurons) if saved_neurons is not None else (int(self.hyper_widgets['Neurons'].get()) if 'Neurons' in self.hyper_widgets else int(MODEL_LIBRARY['ANN (MLP)']['Neurons']['default']))
        activation = saved_activation if saved_activation is not None else (self.hyper_widgets['Activation'].get() if 'Activation' in self.hyper_widgets else MODEL_LIBRARY['ANN (MLP)']['Activation']['options'][0])
        model = MLP(
            in_dim=saved_input_dim,
            hidden_layers=hidden_layers,
            neurons=neurons,
            activation=activation,
            num_outputs=saved_num_outputs
        )
        model.load_state_dict(state_dict)
        model.eval()
        return model, saved_num_outputs
    def _predict_with_model(self, model, X, target_type):
        if hasattr(model, 'predict'):
            return model.predict(X)
        if isinstance(model, nn.Module):
            with torch.no_grad():
                t_data = torch.tensor(X, dtype=torch.float32)
                raw = model(t_data)
                if target_type == "Regression":
                    return raw.cpu().numpy().flatten()
                return torch.argmax(raw, dim=1).cpu().numpy()
        return None
    def execute(self):
        if self.df is None:
            messagebox.showerror("Error", "No dataset loaded. Please load a CSV or Excel file first.")
            return
        if self.output_col is None or not self.input_cols:
            messagebox.showerror("Error", "Select input and target columns before training.")
            return
        if self.output_col not in self.df.columns:
            messagebox.showerror("Error", "Selected target column is not present in the dataset.")
            return
        if self.partition_mode_var.get() == "external" and self.external_test_df is None:
            messagebox.showerror("Error", "Load an external test dataset before training.")
            return
        print(">>> Starting the pipeline")
        self.cancel_event.clear()
        self.epoch_history = {"train": [], "val": [], "test": []}
        self.set_training_state(True)
        feats = {k: v.get() for k, v in self.feats.items()}
        params = {k: w.get() for k, w in self.hyper_widgets.items()}
        target_type = self.target_type_var.get()
        model_name = self.model_choice.get()
        split_pct = int(self.split_slider.get())
        t = threading.Thread(
            target=self._run_training, 
            daemon=True, 
            args=(feats, params, target_type, model_name, split_pct)
        )
        self.training_thread = t
        t.start()
 # --- WORKER THREAD ---
    def _run_training(self, feats, params, target_type, selected_model_name, train_pct):
        try:
            device = "cuda" if torch.cuda.is_available() else "cpu"
            print(">>> Starting the training pipeline...")
            print("\n" + "="*30)
            print(f"MODEL SUMMARY")
            print(f"Method: {selected_model_name}")
            print(f"Task:   {target_type}")
            print("Hyperparameters:")
            summary_lines = [
                f"Method: {selected_model_name}",
                f"Task:   {target_type}",
                "Hyperparameters:"
            ]
            for k, v in params.items():
                print(f"  - {k}: {v}")
                summary_lines.append(f"  - {k}: {v}")
            self.last_model_summary = "\n".join(summary_lines)
            self.last_results = {}
            print("="*30 + "\n")
 # --- 1. DATA PREP --
            external_mode = self.partition_mode_var.get() == "external"
            if selected_model_name in ["LSTM", "CNN"]:                  # CNN and LSTM need 3D (Windows, Time, Features)
                win_size = int(self.window_size_entry.get())
                step = int(win_size * (1 - int(self.overlap_slider.get()) / 100))
                mode = self.extract_mode_var.get()
                X_seq, y_seq = [], []
                for start in range(0, len(self.df) - win_size + 1, step):
                    block = self.df[self.input_cols].iloc[start : start + win_size].astype(float).values
                    if mode == "sliding":
                        X_seq.append(block)
                    elif mode == "fourier":
                        fft_block = np.abs(np.fft.rfft(block - np.mean(block, axis=0), axis=0))
                        X_seq.append(fft_block.T)
                    else:
                        fft_block = np.log1p(np.abs(np.fft.rfft(block - np.mean(block, axis=0), axis=0)))
                        X_seq.append(fft_block.T)
                    y_seq.append(self.df[self.output_col].iloc[start + win_size - 1])
                X, y = np.array(X_seq), np.array(y_seq)
            else:                                                        # MLP/SVM/RBF need 2D (Flat features)
                if self.extract_var.get():
                    mode = self.extract_mode_var.get()
                    if mode == "sliding":
                        X, y = sliding_window_features(self.df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats)
                    elif mode == "fourier":
                        X, y = fourier_window_features(self.df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats)
                    else:
                        X, y = frequency_domain_features(self.df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats)
                else:
                    X, y = self.df[self.input_cols].values, self.df[self.output_col].values      
            if external_mode:
                if self.external_test_df is None:
                    raise ValueError("External test dataset not loaded.")
                if selected_model_name in ["LSTM", "CNN"]:
                    win_size = int(self.window_size_entry.get())
                    step = int(win_size * (1 - int(self.overlap_slider.get()) / 100))
                    mode = self.extract_mode_var.get()
                    X_ext_seq, y_ext_seq = [], []
                    for start in range(0, len(self.external_test_df) - win_size + 1, step):
                        block = self.external_test_df[self.input_cols].iloc[start : start + win_size].astype(float).values
                        if mode == "sliding":
                            X_ext_seq.append(block)
                        elif mode == "fourier":
                            fft_block = np.abs(np.fft.rfft(block - np.mean(block, axis=0), axis=0))
                            X_ext_seq.append(fft_block.T)
                        else:
                            fft_block = np.log1p(np.abs(np.fft.rfft(block - np.mean(block, axis=0), axis=0)))
                            X_ext_seq.append(fft_block.T)
                        y_ext_seq.append(self.external_test_df[self.output_col].iloc[start + win_size - 1])
                    X_test_external, y_test_external = np.array(X_ext_seq), np.array(y_ext_seq)
                else:
                    if self.extract_var.get():
                        X_test_external, y_test_external = sliding_window_features(self.external_test_df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats)
                    else:
                        X_test_external, y_test_external = self.external_test_df[self.input_cols].values, self.external_test_df[self.output_col].values
            y_working = y.copy()  # Encoding and Splitting
            if target_type != "Regression":
                classes, y_working = np.unique(y, return_inverse=True)
                if external_mode:
                    y_test_external = np.array([classes.tolist().index(val) if val in classes else -1 for val in y_test_external])
                    if np.any(y_test_external == -1):
                        raise ValueError("External test dataset contains unseen class labels.")
            num_classes = len(classes) if target_type != "Regression" else 0
            if target_type == "Regression":
                num_out = 1
            elif target_type == "Binary":
                if num_classes != 2:
                    self.after(0, lambda: messagebox.showerror("Error", f"Binary classification selected, but target has {num_classes} classes. Please select Multi-Class or ensure exactly 2 classes."))
                    return
                num_out = 2
            else:  # Multi-Class
                if num_classes < 2:
                    self.after(0, lambda: messagebox.showerror("Error", "Target must have at least 2 classes for classification."))
                    return
                num_out = num_classes
            if external_mode:
                Xtr, ytr = X, y_working
                Xv, yv = None, None
                Xt, yt = X_test_external, y_test_external
            else:
                Xtr, ytr, Xv, yv, Xt, yt = split_data(X, y_working, train_pct)
            Xtr, Xv, Xt = self.apply_normalization_to_sets(Xtr, Xv, Xt)
            if len(Xtr) == 0 or len(ytr) == 0:
                raise ValueError("Training split is empty after partitioning.")
            if Xv is not None and len(Xv) == 0:
                Xv, yv = None, None
            if Xt is not None and len(Xt) == 0:
                Xt, yt = None, None
 # --- 2. MODEL EXECUTION ---
            def progress_cb(current, total, text=None):
                self.after(0, lambda current=current, total=total, text=text: self.update_progress(current, total, text))

            if selected_model_name == "LSTM":
                model = train_lstm(Xtr, ytr, Xv, yv, params, num_out, target_type, self.cancel_event, self.epoch_history, Xt, yt, progress_callback=progress_cb)
            elif selected_model_name == "CNN":
                model = train_cnn(Xtr, ytr, Xv, yv, params, num_out, target_type, self.cancel_event, self.epoch_history, Xt, yt, progress_callback=progress_cb)               
            elif selected_model_name in ["SVM", "LR (Linear/Logistic)", "Decision Tree", "Random Forest", "Gradient Boosting"]:  # Combined block for Scikit-learn models
                if selected_model_name == "SVM":
                    model = train_svm(Xtr, ytr, params, target_type, progress_callback=progress_cb)
                elif selected_model_name == "LR (Linear/Logistic)":
                    model = train_lr(Xtr, ytr, params, target_type, progress_callback=progress_cb)
                else:
                    model = train_tree_methods(Xtr, ytr, params, target_type, selected_model_name, progress_callback=progress_cb)
            elif selected_model_name == "RBF Network":
                model = train_rbf(Xtr, ytr, Xv, yv, params, num_out, target_type, self.cancel_event, self.epoch_history, Xt, yt, progress_callback=progress_cb)
            else: # Default to MLP
                model = train_mlp(Xtr, ytr, Xv, yv, params, num_out, target_type, self.cancel_event, self.epoch_history, Xt, yt, progress_callback=progress_cb)
 # --- 3. POST-TRAINING PREDICTIONS (All Models) ---
            self.trained_model = model 
            if self.cancel_event.is_set():
                print(">> Training was canceled before completion.")
                self.after(0, lambda: self.status_label.configure(text="Training canceled by user."))
                return
            if selected_model_name in ["SVM", "LR (Linear/Logistic)", "Decision Tree", "Random Forest", "Gradient Boosting"]:
                train_preds = model.predict(Xtr)
                val_preds   = model.predict(Xv) if Xv is not None else None
                test_preds  = model.predict(Xt)
            else:
                model.eval()
                with torch.no_grad():
                    def get_nn_preds(data):
                        if data is None or len(data) == 0:
                            return np.array([])
                        t_data = torch.tensor(data, dtype=torch.float32).to(device)
                        raw = model(t_data)
                        if target_type == "Regression":
                            return raw.detach().cpu().numpy().reshape(-1)
                        else:
                            return torch.argmax(raw, dim=1).cpu().numpy()
                    train_preds = get_nn_preds(Xtr)
                    val_preds   = get_nn_preds(Xv) if Xv is not None else None
                    test_preds  = get_nn_preds(Xt)
            if target_type != "Regression":
                print(f"Unique test predictions (encoded): {np.unique(test_preds)}")
                test_preds_orig = classes[test_preds.astype(int)]
                print(f"Unique test predictions (original): {np.unique(test_preds_orig)}")
            self.print_detailed_metrics(ytr, train_preds, "TRAINING", target_type)
            if Xv is not None and yv is not None and len(Xv) > 0 and val_preds is not None and len(val_preds) > 0:
                self.print_detailed_metrics(yv, val_preds, "VALIDATION", target_type)
            if Xt is not None and yt is not None and len(Xt) > 0 and test_preds is not None and len(test_preds) > 0:
                self.print_detailed_metrics(yt, test_preds, "TESTING", target_type)
                self.after(0, lambda: self.update_post_training_ui(yt, test_preds))
            elif Xt is None or len(Xt) == 0:
                self.after(0, lambda: self.metrics_label.configure(text="Training complete. No test set available for evaluation."))
        except Exception as e:
            err_msg = str(e)
            print(f">> Error during training: {err_msg}")
            self.after(0, lambda err_msg=err_msg: messagebox.showerror("Training Error", err_msg))
        finally:
            self.after(0, lambda: self.set_training_state(False))
    def perform_grid_search(self):
        if self.df is None:
            messagebox.showerror("Error", "No dataset loaded. Please load a CSV or Excel file first.")
            return
        if self.output_col is None or not self.input_cols:
            messagebox.showerror("Error", "Select input and target columns before running grid search.")
            return
        if self.partition_mode_var.get() == "external" and self.external_test_df is None:
            messagebox.showerror("Error", "Load an external test dataset before running grid search.")
            return
        self.cancel_event.clear()
        self.set_training_state(True)
        params = {k: w.get() for k, w in self.hyper_widgets.items()}
        model_name = self.model_choice.get()
        target_type = self.target_type_var.get()
        split_pct = int(self.split_slider.get())
        threading.Thread(
            target=self._run_grid_search,
            daemon=True,
            args=(model_name, target_type, params, split_pct)
        ).start()
    def _run_grid_search(self, model_name, target_type, base_params, train_pct):
        try:
            print(">> Starting grid search...")
            candidate_grid = self._build_grid_candidates(model_name, base_params)
            if not candidate_grid:
                raise ValueError("No grid search candidates could be generated for this model.")
            best_score = None
            best_params = None
            iterations = []
            self.after(0, lambda total=len(candidate_grid): self.update_progress(0, total, "Grid search starting..."))
            for idx, params in enumerate(candidate_grid, start=1):
                if self.cancel_event.is_set():
                    print(">> Grid search canceled.")
                    break
                print(f">> Grid Search candidate {idx}/{len(candidate_grid)}: {params}")
                result = self._evaluate_parameters(model_name, target_type, params, train_pct)
                if not self.cancel_event.is_set():
                    self.after(0, lambda idx=idx, total=len(candidate_grid): self.update_progress(idx, total, f"Grid Search {idx}/{total}"))
                score = result["score"]
                iteration_info = {
                    "iteration": idx,
                    "params": params.copy(),
                    "score": score,
                    "accuracy": result.get("accuracy"),
                    "error": result.get("error")
                }
                iterations.append(iteration_info)
                print(f"   Candidate score: {score:.4f}, accuracy: {iteration_info['accuracy']}, error: {iteration_info['error']}")
                if best_score is None or score > best_score:
                    best_score = score
                    best_params = params.copy()
            if best_params is None:
                raise ValueError("Grid search did not complete any candidate evaluations.")
            self.grid_search_results = {
                "model": model_name,
                "target_type": target_type,
                "params": best_params,
                "score": best_score,
                "iterations": iterations
            }
            self.after(0, lambda: self._apply_grid_search_results(best_params, target_type, best_score))
        except Exception as e:
            err_msg = str(e)
            print(f">> Grid search error: {err_msg}")
            self.after(0, lambda err_msg=err_msg: messagebox.showerror("Grid Search Error", err_msg))
        finally:
            self.after(0, lambda: self.set_training_state(False))
    def _build_grid_candidates(self, model_name, current_params):
        if model_name not in MODEL_LIBRARY:
            return []
        param_values = {}
        for name, info in MODEL_LIBRARY[model_name].items():
            if info["type"] == "list":
                param_values[name] = info["options"]
            else:
                try:
                    value = float(current_params.get(name, info.get("default", 1)))
                except Exception:
                    value = float(info.get("default", 1))
                if info["type"] == "int":
                    value = max(1, int(value))
                    param_values[name] = sorted(set([value, max(1, value // 2), max(1, value * 2)]))
                else:
                    param_values[name] = sorted(set([value, max(value / 10, 1e-6), value * 10]))
        sorted_items = sorted(param_values.items(), key=lambda item: len(item[1]), reverse=True)
        if len(sorted_items) > 3:
            sorted_items = sorted_items[:3]
        names, value_lists = zip(*sorted_items) if sorted_items else ([], [])
        grid = []
        for combo in product(*value_lists):
            candidate = current_params.copy()
            for name, value in zip(names, combo):
                candidate[name] = str(int(value)) if MODEL_LIBRARY[model_name][name]["type"] == "int" else str(value)
            grid.append(candidate)
        return grid
    def _evaluate_parameters(self, model_name, target_type, params, train_pct):
        external_mode = self.partition_mode_var.get() == "external"
        if model_name in ["LSTM", "CNN"]:
            win_size = int(self.window_size_entry.get())
            step = int(win_size * (1 - int(self.overlap_slider.get()) / 100))
            X_seq, y_seq = [], []
            for start in range(0, len(self.df) - win_size + 1, step):
                X_seq.append(self.df[self.input_cols].iloc[start : start + win_size].values)
                y_seq.append(self.df[self.output_col].iloc[start + win_size - 1])
            X, y = np.array(X_seq), np.array(y_seq)
        else:
            if self.extract_var.get():
                feats = {k: v.get() for k, v in self.feats.items()}
                X, y = sliding_window_features(self.df, self.input_cols, self.output_col, int(self.window_size_entry.get()), int(self.overlap_slider.get()), feats)
            else:
                X, y = self.df[self.input_cols].values, self.df[self.output_col].values
        y_working = y.copy()
        classes = None
        if target_type != "Regression":
            classes, y_working = np.unique(y, return_inverse=True)
        if external_mode:
            if self.external_test_df is None:
                raise ValueError("External test dataset not loaded.")
            Xtr, ytr = X, y_working
            Xv, yv = None, None
            Xt = self.external_test_df[self.input_cols].values
            yt_orig = self.external_test_df[self.output_col].values
            if target_type != "Regression":
                yt = np.array([classes.tolist().index(val) if val in classes else -1 for val in yt_orig])
                if np.any(yt == -1):
                    raise ValueError("External test dataset contains unseen class labels.")
            else:
                yt = yt_orig
        else:
            Xtr, ytr, Xv, yv, Xt, yt = split_data(X, y_working, train_pct)
        if model_name == "LSTM":
            model = train_lstm(Xtr, ytr, Xv, yv, params, target_type, self.cancel_event, None, Xt, yt)
        elif model_name == "CNN":
            model = train_cnn(Xtr, ytr, Xv, yv, params, target_type, self.cancel_event, None, Xt, yt)
        elif model_name in ["SVM", "LR (Linear/Logistic)", "Decision Tree", "Random Forest", "Gradient Boosting"]:
            if model_name == "SVM":
                model = train_svm(Xtr, ytr, params, target_type)
            elif model_name == "LR (Linear/Logistic)":
                model = train_lr(Xtr, ytr, params, target_type)
            else:
                model = train_tree_methods(Xtr, ytr, params, target_type, model_name)
        elif model_name == "RBF Network":
            model = train_rbf(Xtr, ytr, Xv, yv, params, target_type, self.cancel_event, None, Xt, yt)
        else:
            model = train_mlp(Xtr, ytr, Xv, yv, params, target_type, self.cancel_event, None, Xt, yt)
        if self.cancel_event.is_set():
            raise ValueError("Grid search was canceled.")
        eval_X = Xt if Xt is not None else Xv if Xv is not None else Xtr
        eval_y = yt if Xt is not None else yv if Xv is not None else ytr
        preds = self._predict_with_model(model, eval_X, target_type)
        if target_type == "Regression":
            mse = mean_squared_error(eval_y, preds)
            return {"score": -mse, "accuracy": None, "error": mse}
        acc = float((preds == eval_y).mean())
        return {"score": acc, "accuracy": acc, "error": 1.0 - acc}
    def _apply_grid_search_results(self, best_params, target_type, best_score):
        for name, widget in self.hyper_widgets.items():
            if name in best_params:
                value = str(best_params[name])
                if hasattr(widget, 'set'):
                    widget.set(value)
                elif hasattr(widget, 'delete') and hasattr(widget, 'insert'):
                    widget.delete(0, 'end')
                    widget.insert(0, value)
                else:
                    try:
                        widget.configure(text=value)
                    except Exception:
                        pass
        current_results = getattr(self, 'grid_search_results', {})
        current_results.update({
            "model": self.model_choice.get(),
            "target_type": target_type,
            "params": best_params,
            "score": best_score
        })
        self.grid_search_results = current_results
        self.status_label.configure(text=f"Grid search completed. Best score: {best_score:.4f}")
        if messagebox.askyesno("Export Grid Search Results", "Grid search is complete. Would you like to save the best results to a Word document?"):
            self.export_grid_search_results_to_word()
        else:
            messagebox.showinfo("Grid Search Complete", f"Best candidate found with score {best_score:.4f}. Hyperparameters updated.")
    def export_grid_search_results_to_word(self):
        if not getattr(self, 'grid_search_results', None):
            messagebox.showerror("Error", "No grid search results available to export.")
            return
        save_path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Document", "*.docx")])
        if not save_path:
            return
        try:
            doc = Document()
            doc.add_heading("Grid Search Results", level=1)
            doc.add_paragraph(f"Model: {self.grid_search_results['model']}")
            doc.add_paragraph(f"Task: {self.grid_search_results['target_type']}")
            doc.add_paragraph(f"Best Score: {self.grid_search_results['score']:.4f}")
            doc.add_heading("Grid Search Iterations", level=2)
            iterations = self.grid_search_results.get('iterations', [])
            if iterations:
                table = doc.add_table(rows=len(iterations) + 1, cols=4)
                table.style = 'Light List Accent 1'
                headers = ["Iteration", "Accuracy", "Error", "Parameters"]
                for col_index, header in enumerate(headers):
                    table.cell(0, col_index).text = header
                for i, entry in enumerate(iterations, start=1):
                    table.cell(i, 0).text = str(entry['iteration'])
                    table.cell(i, 1).text = f"{entry['accuracy']:.4f}" if entry['accuracy'] is not None else "N/A"
                    table.cell(i, 2).text = f"{entry['error']:.4f}" if entry['error'] is not None else "N/A"
                    params_text = ", ".join(f"{k}={v}" for k, v in entry['params'].items())
                    table.cell(i, 3).text = params_text
            else:
                doc.add_paragraph("No iteration details were recorded.")
            doc.add_heading("Best Hyperparameters", level=2)
            for name, value in self.grid_search_results['params'].items():
                doc.add_paragraph(f"{name}: {value}")
            doc.save(save_path)
            messagebox.showinfo("Export Complete", f"Grid search results exported to {save_path}")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))
    def export_results_to_word(self):
        if not self.last_results:
            messagebox.showerror("Error", "No training results available to export.")
            return
        if not self.last_model_summary:
            messagebox.showerror("Error", "No model summary available to export.")
            return
        save_path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Document", "*.docx")])
        if not save_path:
            return
        try:
            doc = Document()
            doc.add_heading("AI Model Training Results", level=1)
            for line in self.last_model_summary.splitlines():
                doc.add_paragraph(line)
            for split in ["TRAINING", "VALIDATION", "TESTING"]:
                if split not in self.last_results:
                    continue
                metrics = self.last_results[split]
                doc.add_heading(f"{split} Metrics", level=2)
                is_regression = "r2" in metrics
                if is_regression:
                    doc.add_paragraph(f"R2: {metrics['r2']:.4f}")
                    doc.add_paragraph(f"MSE: {metrics['mse']:.4f}")
                    doc.add_paragraph(f"RMSE: {metrics['rmse']:.4f}")
                    doc.add_paragraph(f"MAE: {metrics['mae']:.4f}")
                else:
                    doc.add_paragraph(f"Accuracy: {metrics['accuracy']*100:.2f}%")
                    doc.add_paragraph(f"Precision (weighted): {metrics['precision']:.4f}")
                    doc.add_paragraph(f"Recall (weighted): {metrics['recall']:.4f}")
                    doc.add_paragraph(f"F1 Score (weighted): {metrics['f1']:.4f}")
                    doc.add_paragraph(f"Support total: {sum(metrics['support'])}")
                    doc.add_heading("Class-level Metrics", level=3)
                    for label, p, r, f, s in zip(metrics['labels'], metrics['per_class_precision'], metrics['per_class_recall'], metrics['per_class_f1'], metrics['support']):
                        doc.add_paragraph(f"Class {label}: Precision: {p:.4f} | Recall: {r:.4f} | F1: {f:.4f} | Support: {s}")
                    if "confusion_matrix" in metrics:
                        doc.add_heading("Confusion Matrix", level=3)
                        cm_data = metrics["confusion_matrix"]
                        if cm_data:
                            rows = len(cm_data)
                            cols = len(cm_data[0])
                            table = doc.add_table(rows=rows + 1, cols=cols + 1)
                            table.style = 'Light List Accent 1'
                            table.cell(0, 0).text = "True\\Pred"
                            for j, label in enumerate(metrics['labels'], start=1):
                                table.cell(0, j).text = str(label)
                            for i, row in enumerate(cm_data, start=1):
                                table.cell(i, 0).text = str(metrics['labels'][i-1])
                                for j, value in enumerate(row, start=1):
                                    table.cell(i, j).text = str(value)
                    else:
                        doc.add_paragraph("No confusion matrix available.")
                doc.add_paragraph("\n")
            if self.epoch_history and self.epoch_history.get("train"):
                target_type = self.target_type_var.get()
                heading = "Loss per Epoch" if target_type == "Regression" else "Accuracy per Epoch"
                doc.add_heading(heading, level=2)
                epochs = list(range(1, len(self.epoch_history["train"]) + 1))
                has_val = any(v is not None for v in self.epoch_history.get("val", []))
                has_test = any(v is not None for v in self.epoch_history.get("test", []))
                cols = ["Epoch", f"Train {'Loss' if target_type == 'Regression' else 'Accuracy'}"]
                if has_val:
                    cols.append(f"Validation {'Loss' if target_type == 'Regression' else 'Accuracy'}")
                if has_test:
                    cols.append(f"Test {'Loss' if target_type == 'Regression' else 'Accuracy'}")
                table = doc.add_table(rows=len(epochs) + 1, cols=len(cols))
                table.style = 'Light List Accent 1'
                for j, col_name in enumerate(cols):
                    table.cell(0, j).text = col_name
                for i, epoch in enumerate(epochs, start=1):
                    table.cell(i, 0).text = str(epoch)
                    train_val = self.epoch_history["train"][i - 1]
                    if target_type == "Regression":
                        table.cell(i, 1).text = f"{train_val:.4f}" if train_val is not None else "N/A"
                    else:
                        table.cell(i, 1).text = f"{train_val*100:.2f}%" if train_val is not None else "N/A"
                    col_index = 2
                    if has_val:
                        val = self.epoch_history["val"][i - 1]
                        if target_type == "Regression":
                            table.cell(i, col_index).text = f"{val:.4f}" if val is not None else "N/A"
                        else:
                            table.cell(i, col_index).text = f"{val*100:.2f}%" if val is not None else "N/A"
                        col_index += 1
                    if has_test:
                        test_val = self.epoch_history["test"][i - 1]
                        if target_type == "Regression":
                            table.cell(i, col_index).text = f"{test_val:.4f}" if test_val is not None else "N/A"
                        else:
                            table.cell(i, col_index).text = f"{test_val*100:.2f}%" if test_val is not None else "N/A"
            doc.add_heading("Sample Ground Truth vs Predictions", level=2)
            for true_val, pred_val in zip(self.last_y_true[:20], self.last_y_pred[:20]):
                doc.add_paragraph(f"True: {true_val} | Pred: {pred_val}")
            doc.save(save_path)
            messagebox.showinfo("Export Complete", f"Results exported to {save_path}")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))
    def plot_results(self):
        fig, ax = plt.subplots(figsize=(8, 5)) 
        if self.target_type_var.get() == "Regression":
            ax.scatter(self.last_y_true, self.last_y_pred, alpha=0.5)
            ax.plot([min(self.last_y_true), max(self.last_y_true)], 
                [min(self.last_y_true), max(self.last_y_true)], 'r--')
            ax.set_title("Regression Accuracy")
        else:
            cm = confusion_matrix(self.last_y_true, self.last_y_pred)
            # 2. Tell the display to use the 'ax' we created above
            disp = ConfusionMatrixDisplay(confusion_matrix=cm)
            disp.plot(cmap='Blues', ax=ax) 
            ax.set_title("Confusion Matrix")
        plt.show()
    def plot_accuracy_per_epoch(self):
        if not self.epoch_history or not self.epoch_history.get("train"):
            messagebox.showerror("Error", "No epoch history available. Train a model with epochs first.")
            return
        target_type = self.target_type_var.get()
        epochs = list(range(1, len(self.epoch_history["train"]) + 1))
        val_series = [float(v) if v is not None else np.nan for v in self.epoch_history["val"]]
        test_series = [float(v) if v is not None else np.nan for v in self.epoch_history["test"]]
        fig, ax = plt.subplots(figsize=(8, 5))
        ax.plot(epochs, self.epoch_history["train"], marker='o', label="Train")
        if any(not np.isnan(v) for v in val_series):
            ax.plot(epochs, val_series, marker='o', label="Validation")
        if any(not np.isnan(v) for v in test_series):
            ax.plot(epochs, test_series, marker='o', label="Test")
        if target_type == "Regression":
            ax.set_title("Loss per Epoch")
            ax.set_ylabel("Loss")
        else:
            ax.set_title("Accuracy per Epoch")
            ax.set_ylabel("Accuracy")
            ax.set_ylim(0.0, 1.0)
        ax.set_xlabel("Epoch")
        ax.legend()
        ax.grid(True)
        plt.show()
    def print_detailed_metrics(self, y_true, y_pred, split_name, target_type):
        from sklearn.metrics import precision_recall_fscore_support, mean_absolute_error
        print(f"\n--- {split_name} Results ---")
        result = {"split": split_name}
        if target_type == "Regression":
            mse = mean_squared_error(y_true, y_pred)
            rmse = np.sqrt(mse)
            mae = mean_absolute_error(y_true, y_pred)
            r2 = r2_score(y_true, y_pred)
            print(f"  R2 Score:  {r2:.4f}\n  MSE:       {mse:.4f}\n  RMSE:      {rmse:.4f}\n  MAE:       {mae:.4f}")
            result.update({"mse": mse, "rmse": rmse, "mae": mae, "r2": r2})
        else:
            acc = (y_true == y_pred).mean()
            precision, recall, f1, _ = precision_recall_fscore_support(y_true, y_pred, average='weighted')
            labels = np.unique(y_true)
            per_precision, per_recall, per_f1, support = precision_recall_fscore_support(y_true, y_pred, average=None, labels=labels)
            cm = confusion_matrix(y_true, y_pred)
            print(f"  Accuracy:  {acc*100:.2f}%\n  Precision: {precision:.4f}\n  Recall:    {recall:.4f}\n  F1-Score:  {f1:.4f}")
            print("  Confusion Matrix:")
            for row in cm:
                print("    " + "  ".join(str(int(x)) for x in row))
            result.update({
                "accuracy": acc,
                "precision": precision,
                "recall": recall,
                "f1": f1,
                "labels": labels.tolist(),
                "per_class_precision": per_precision.tolist(),
                "per_class_recall": per_recall.tolist(),
                "per_class_f1": per_f1.tolist(),
                "support": support.tolist(),
                "confusion_matrix": cm.tolist()
            })
        self.last_results[split_name] = result
if __name__ == "__main__":
    try:
        app = DataAnalyzerGUI()
        app.mainloop()
    except Exception as e:
        import traceback
        error_msg = f"Error starting application:\n{str(e)}\n\n{traceback.format_exc()}"
        print(error_msg)
        # Write to error log file (use UTF-8 to support all characters)
        try:
            with open("error_log.txt", "w", encoding="utf-8") as f:
                f.write(error_msg)
        except Exception:
            # Fallback to default encoding if UTF-8 write fails for any reason
            with open("error_log.txt", "w", errors="ignore") as f:
                f.write(error_msg)
        try:
            messagebox.showerror("Application Error", error_msg)
        except:
            pass
